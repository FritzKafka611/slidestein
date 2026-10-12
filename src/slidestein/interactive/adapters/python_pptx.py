"""PythonPptxPowerPointHostAdapter — python-pptx backed PPTX capture and patching.

Transactional safety guarantee (M11.1 Req 2):
  Verification runs on the TEMP file BEFORE os.replace.
  The original PPTX is never overwritten until full verification passes.
  On verification failure: temp deleted, original byte-for-byte unchanged.

z_order, editable_text, parent_group_path are intentionally excluded from the
canonical state hash (M11.1 Req 17):
  - z_order: derived from shape iteration order; unreliable across group nesting levels
  - editable_text: a derived flag (True iff shape has text_frame and is not a group)
  - parent_group_path: derivable from shape_path by string split
These fields are captured for convenience but do not affect the hash.
"""

from __future__ import annotations

import hashlib
import os
import uuid
from datetime import datetime, timezone
from pathlib import Path
from typing import Optional

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from slidestein.interactive.hashing import compute_slide_state_hash
from slidestein.interactive.models import (
    PatchExecutionResult,
    ShapeState,
    SlideState,
    SlidePatchOperation,
    SlidePatchPlan,
)
from slidestein.structural.diagnosis import _build_path_map, _slide_space_geometry


class PythonPptxPowerPointHostAdapter:
    """Captures and patches PPTX slide state using python-pptx file I/O."""

    def capture_slide_state(
        self,
        pptx_path: Path,
        slide_number: int,
        *,
        deck_id: Optional[str] = None,
        slide_id: Optional[str] = None,
    ) -> SlideState:
        prs = Presentation(str(pptx_path))
        slide = prs.slides[slide_number - 1]
        path_map = _build_path_map(slide.shapes)

        slide_width = int(prs.slide_width)
        slide_height = int(prs.slide_height)

        shapes: list[ShapeState] = []
        path_list = list(path_map.keys())
        for z_order, shape_path in enumerate(path_list):
            shape = path_map[shape_path]
            geom = _slide_space_geometry(path_map, shape_path)
            if geom is None:
                continue

            shape_type = _detect_shape_type(shape)
            is_group = shape_type == "group"
            parent_group_path = _parent_path(shape_path)

            text: Optional[str] = None
            editable_text = False
            if not is_group and hasattr(shape, "text_frame"):
                try:
                    text = shape.text_frame.text
                    editable_text = True
                except Exception:
                    pass

            fingerprint = _compute_formatting_fingerprint(shape)
            rotation: Optional[float] = None
            try:
                rotation = float(shape.rotation) if hasattr(shape, "rotation") else None
            except Exception:
                pass

            shapes.append(
                ShapeState(
                    shape_id=shape.shape_id,
                    shape_path=shape_path,
                    name=shape.name,
                    shape_type=shape_type,
                    parent_group_path=parent_group_path,
                    text=text,
                    x=geom["x"],
                    y=geom["y"],
                    width=geom["width"],
                    height=geom["height"],
                    rotation=rotation,
                    z_order=z_order,
                    editable_text=editable_text,
                    formatting_fingerprint=fingerprint,
                )
            )

        state_hash = compute_slide_state_hash(shapes, slide_width, slide_height)
        captured_at = datetime.now(timezone.utc).isoformat()

        return SlideState(
            deck_id=deck_id,
            slide_id=slide_id,
            slide_number=slide_number,
            slide_width=slide_width,
            slide_height=slide_height,
            shapes=shapes,
            state_hash=state_hash,
            captured_at=captured_at,
        )

    def apply_patch(
        self,
        pptx_path: Path,
        patch: SlidePatchPlan,
    ) -> PatchExecutionResult:
        # ── Step 1: Capture authoritative before state ─────────────────────
        before_state = self.capture_slide_state(pptx_path, patch.slide_number)

        if before_state.state_hash != patch.base_state_hash:
            return PatchExecutionResult(
                status="rejected",
                rejection_reason="stale_base_state",
                base_state_hash=patch.base_state_hash,
                observed_pre_state_hash=before_state.state_hash,
                before_state=before_state,
            )

        # ── Step 2: Validate every operation before any file I/O ───────────
        shape_map = {s.shape_path: s for s in before_state.shapes}

        for op in patch.operations:
            if op.shape_path not in shape_map:
                return PatchExecutionResult(
                    status="rejected",
                    rejection_reason="missing_target_shape",
                    base_state_hash=patch.base_state_hash,
                    observed_pre_state_hash=before_state.state_hash,
                    before_state=before_state,
                    skipped_operations=list(patch.operations),
                )

            target = shape_map[op.shape_path]

            if target.shape_id != op.shape_id:
                return PatchExecutionResult(
                    status="rejected",
                    rejection_reason="shape_identity_mismatch",
                    base_state_hash=patch.base_state_hash,
                    observed_pre_state_hash=before_state.state_hash,
                    before_state=before_state,
                    skipped_operations=list(patch.operations),
                )

            if not target.editable_text:
                return PatchExecutionResult(
                    status="rejected",
                    rejection_reason="target_not_editable",
                    base_state_hash=patch.base_state_hash,
                    observed_pre_state_hash=before_state.state_hash,
                    before_state=before_state,
                    skipped_operations=list(patch.operations),
                )

            if op.expected_before is not None and target.text != op.expected_before:
                return PatchExecutionResult(
                    status="rejected",
                    rejection_reason="precondition_failed",
                    base_state_hash=patch.base_state_hash,
                    observed_pre_state_hash=before_state.state_hash,
                    before_state=before_state,
                    skipped_operations=list(patch.operations),
                )

        # ── Step 3: Apply mutations in memory only — no writes yet ─────────
        prs = Presentation(str(pptx_path))
        slide = prs.slides[patch.slide_number - 1]
        path_map_writable = _build_path_map(slide.shapes)
        self._apply_mutations(slide, path_map_writable, patch.operations)

        # ── Step 4: Save to TEMP file ───────────────────────────────────────
        temp = pptx_path.parent / f"._m11tmp_{uuid.uuid4().hex[:12]}_{pptx_path.name}"
        try:
            prs.save(str(temp))
        except Exception:
            try:
                temp.unlink()
            except Exception:
                pass
            raise

        # ── Step 5: Capture candidate state FROM TEMP (original untouched) ─
        try:
            candidate_state = self.capture_slide_state(temp, patch.slide_number)
        except Exception:
            try:
                temp.unlink()
            except Exception:
                pass
            raise

        # ── Step 6: Verify candidate state completely ───────────────────────
        authorized_keys = {f"{op.shape_path}.{op.field}" for op in patch.operations}
        unexpected = _find_unauthorized_changes(before_state, candidate_state, authorized_keys)
        missing_effects = _check_requested_effects(candidate_state, patch.operations)

        if unexpected or missing_effects:
            try:
                temp.unlink()
            except Exception:
                pass
            return PatchExecutionResult(
                status="rejected",
                rejection_reason="verification_failed",
                base_state_hash=patch.base_state_hash,
                observed_pre_state_hash=before_state.state_hash,
                before_state=before_state,
                verification_passed=False,
                unexpected_changed_fields=unexpected + missing_effects,
            )

        # ── Step 7: Commit — original first touched HERE, only after verification ──
        try:
            os.replace(str(temp), str(pptx_path))
        except Exception:
            try:
                temp.unlink()
            except Exception:
                pass
            raise

        # ── Step 8: Final capture from original for authoritative after_state ──
        final_state = self.capture_slide_state(pptx_path, patch.slide_number)

        return PatchExecutionResult(
            status="applied",
            base_state_hash=patch.base_state_hash,
            observed_pre_state_hash=before_state.state_hash,
            resulting_state_hash=final_state.state_hash,
            applied_operations=list(patch.operations),
            before_state=before_state,
            after_state=final_state,
            verification_passed=True,
        )

    def _apply_mutations(
        self,
        slide,
        path_map: dict,
        operations: list[SlidePatchOperation],
    ) -> None:
        """Apply text mutations to in-memory shapes. Override in subclasses for testing."""
        for op in operations:
            shape = path_map.get(op.shape_path)
            if shape is None or not hasattr(shape, "text_frame"):
                continue
            new_text = "" if op.operation == "clear_text" else (op.value or "")
            _set_text_preserving_format(shape.text_frame, new_text)


def _find_unauthorized_changes(
    before_state: SlideState,
    after_state: SlideState,
    authorized_keys: set[str],
) -> list[str]:
    """Return field keys that changed without authorization."""
    unexpected: list[str] = []
    before_index = {s.shape_path: s for s in before_state.shapes}
    after_index = {s.shape_path: s for s in after_state.shapes}

    for path in set(before_index) | set(after_index):
        b = before_index.get(path)
        a = after_index.get(path)
        if b is None or a is None:
            key = f"{path}._presence"
            if key not in authorized_keys:
                unexpected.append(key)
            continue
        for field in ("text", "x", "y", "width", "height", "rotation", "formatting_fingerprint"):
            if getattr(b, field) != getattr(a, field):
                key = f"{path}.{field}"
                if key not in authorized_keys:
                    unexpected.append(key)
    return unexpected


def _check_requested_effects(
    candidate_state: SlideState,
    operations: list[SlidePatchOperation],
) -> list[str]:
    """Return identifiers for operations whose expected effect was not observed."""
    missing: list[str] = []
    shape_map = {s.shape_path: s for s in candidate_state.shapes}

    for op in operations:
        key = f"{op.shape_path}.{op.field}"
        target = shape_map.get(op.shape_path)
        if target is None:
            missing.append(f"{key}:shape_missing_after_apply")
            continue
        if target.shape_id != op.shape_id:
            missing.append(f"{key}:shape_id_mismatch_after_apply")
            continue
        expected_text = "" if op.operation == "clear_text" else (op.value or "")
        if target.text != expected_text:
            missing.append(f"{key}:expected={expected_text!r}_got={target.text!r}")
    return missing


def _detect_shape_type(shape) -> str:
    try:
        st = shape.shape_type
        if st == MSO_SHAPE_TYPE.GROUP:
            return "group"
        if st == MSO_SHAPE_TYPE.TEXT_BOX:
            return "text_box"
        if st == MSO_SHAPE_TYPE.AUTO_SHAPE:
            return "auto_shape"
        if st == MSO_SHAPE_TYPE.PICTURE:
            return "picture"
    except Exception:
        pass
    return "other"


def _parent_path(shape_path: str) -> Optional[str]:
    parts = shape_path.split("/")
    return "/".join(parts[:-1]) if len(parts) > 1 else None


def _compute_formatting_fingerprint(shape) -> Optional[str]:
    """Short fingerprint over first-para/first-run font properties only.

    Reads only para[0] alignment and run[0] font (name, size, bold, italic)
    to avoid sensitivity to run-count changes after _set_text_preserving_format.
    Returns None on any exception or missing text_frame.
    """
    try:
        if not hasattr(shape, "text_frame"):
            return None
        tf = shape.text_frame
        paras = tf.paragraphs
        if not paras:
            return None
        para = paras[0]
        alignment = str(para.alignment) if para.alignment is not None else "NONE"
        runs = para.runs
        if not runs:
            canonical = f"align={alignment}"
        else:
            run = runs[0]
            font = run.font
            canonical = (
                f"align={alignment}"
                f"|name={font.name or 'NONE'}"
                f"|size={font.size or 'NONE'}"
                f"|bold={font.bold}"
                f"|italic={font.italic}"
            )
        return hashlib.sha256(canonical.encode("utf-8")).hexdigest()[:16]
    except Exception:
        return None


def _set_text_preserving_format(tf, new_text: str) -> None:
    """Replace text content while keeping run-level formatting.

    Sets first run of first paragraph to new_text; clears extra runs and
    extra paragraphs. Falls back to tf.text assignment if no runs exist.
    """
    try:
        paras = list(tf.paragraphs)
        if not paras:
            tf.text = new_text
            return
        first_para = paras[0]
        runs = list(first_para.runs)
        if runs:
            runs[0].text = new_text
            for r in runs[1:]:
                r.text = ""
        else:
            tf.text = new_text
            return
        for para in paras[1:]:
            for r in para.runs:
                r.text = ""
    except Exception:
        tf.text = new_text
