"""IngestionService — indexes a PPTX deck into the slide library.

For each slide it:
  - extracts the native PowerPoint slide ID from OOXML
  - resolves (or creates) a stable deck_id via lineage rules
  - assigns a UUID5 slide_id stable across edits, reorders, and renames
  - computes content and structure fingerprints
  - extracts available text and structural facts via python-pptx
  - optionally renders a PNG preview via the PPTMasterAdapter
  - persists a SlideRecord to SQLite

Consulting semantics (archetype, communication_jobs, etc.) are left null at
this stage and populated by the classification milestone later.

Dependency rule: imports from slidestein.domain, slidestein.library,
slidestein.identity, and slidestein.pptx.  Never imports consulting logic.
"""

from __future__ import annotations

import hashlib
import logging
import uuid
import zipfile
from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from slidestein.domain.models import SlideRecord, make_slide_id
from slidestein.identity.slide_identity import (
    compute_content_fingerprint,
    compute_deck_fingerprint,
    compute_structure_fingerprint,
    extract_slide_identities,
    make_stable_slide_id,
)
from slidestein.library.store import SlideLibrary
from slidestein.pptx.adapter import PPTMasterAdapter

_LOG = logging.getLogger(__name__)


class IngestionService:
    """Orchestrates single-deck ingestion into the slide library."""

    def __init__(
        self,
        library: SlideLibrary,
        adapter: PPTMasterAdapter,
        previews_dir: Path,
    ) -> None:
        self._library = library
        self._adapter = adapter
        self._previews_dir = previews_dir

    def ingest_deck(
        self,
        pptx_path: Path,
        render_previews: bool = True,
    ) -> list[SlideRecord]:
        """Index every slide in pptx_path and return the resulting SlideRecords.

        Parameters
        ----------
        pptx_path:
            Path to the source PPTX file.  Resolved to an absolute path before use.
        render_previews:
            If True, call adapter.render_preview() for each slide.  Any exception
            from the adapter propagates — use --no-preview (render_previews=False)
            when no rendering backend is available.

        Raises
        ------
        FileNotFoundError
            If pptx_path does not exist.
        ValueError
            If pptx_path is not a .pptx file or cannot be opened by python-pptx.
        """
        pptx_path = pptx_path.resolve()

        if not pptx_path.exists():
            raise FileNotFoundError(f"PPTX file not found: {pptx_path}")
        if not pptx_path.is_file() or pptx_path.suffix.lower() != ".pptx":
            raise ValueError(f"Expected a .pptx file, got: {pptx_path}")

        deck_fingerprint = compute_deck_fingerprint(pptx_path)
        _LOG.info("Ingesting %s  [%s…]", pptx_path.name, deck_fingerprint[:12])

        # Resolve stable deck_id (lineage rules A–D).
        deck_id = _resolve_deck_id(self._library, pptx_path, deck_fingerprint)

        # Mark all current active slides for this deck inactive before re-inserting.
        self._library.mark_slides_inactive(deck_fingerprint=deck_fingerprint)

        # Extract native slide identities from OOXML.
        try:
            slide_identities = extract_slide_identities(pptx_path)
        except (ValueError, KeyError, zipfile.BadZipFile) as exc:
            _LOG.warning("Could not extract native slide IDs: %s", exc)
            slide_identities = []

        try:
            prs = Presentation(str(pptx_path))
        except Exception as exc:
            raise ValueError(f"Cannot open PPTX file {pptx_path}: {exc}") from exc

        slide_width_emu = int(prs.slide_width)
        slide_height_emu = int(prs.slide_height)

        # Build a map of slide_number → SlideIdentity for fast lookup.
        identity_map = {si.slide_number: si for si in slide_identities}

        records: list[SlideRecord] = []

        for idx, slide in enumerate(prs.slides):
            slide_number = idx + 1
            identity = identity_map.get(slide_number)

            if identity is not None:
                slide_id = make_stable_slide_id(deck_id, identity.native_slide_id)
                native_slide_id = identity.native_slide_id
                content_fp = compute_content_fingerprint(pptx_path, identity.slide_part)
                structure_fp = compute_structure_fingerprint(pptx_path, slide_number)
            else:
                # Fallback to legacy ID when OOXML extraction fails.
                slide_id = make_slide_id(deck_fingerprint, slide_number)
                native_slide_id = None
                content_fp = None
                structure_fp = None

            facts = _extract_structural_facts(slide)

            preview_path: Path | None = None
            if render_previews:
                pp = self._previews_dir / deck_fingerprint[:12] / f"{slide_number:03d}.png"
                self._adapter.render_preview(pptx_path, slide_number, pp)
                preview_path = pp

            record = SlideRecord(
                slide_id=slide_id,
                deck_fingerprint=deck_fingerprint,
                source_deck_path=pptx_path,
                slide_number=slide_number,
                extracted_text=facts["extracted_text"],
                slide_width_emu=slide_width_emu,
                slide_height_emu=slide_height_emu,
                text_object_count=facts["text_object_count"],
                image_count=facts["image_count"],
                chart_count=facts["chart_count"],
                table_count=facts["table_count"],
                shape_count=facts["shape_count"],
                preview_path=preview_path,
                deck_id=deck_id,
                native_slide_id=native_slide_id,
                content_fingerprint=content_fp,
                structure_fingerprint=structure_fp,
                is_active=True,
            )

            self._library.upsert_record(record)
            records.append(record)
            _LOG.debug("  slide %d/%d  [%s]", slide_number, len(prs.slides), slide_id)

        _LOG.info("Ingested %d slide(s) from %s", len(records), pptx_path.name)
        return records


# ---------------------------------------------------------------------------
# Deck lineage resolution
# ---------------------------------------------------------------------------


def _resolve_deck_id(
    library: SlideLibrary, pptx_path: Path, deck_fingerprint: str
) -> str:
    """Return the deck_id for this file, creating one if needed.

    Rules (applied in priority order):
      A: Same path as existing deck AND same fingerprint → preserve deck_id unchanged
      B: Same path but different fingerprint → preserve deck_id, update fingerprint
      C: Same fingerprint but different path → preserve deck_id, update path
      D: Neither matches → new UUID4 deck_id
    """
    path_str = str(pptx_path)

    by_path = library.get_deck_by_path(path_str)
    if by_path is not None:
        deck_id = by_path["deck_id"]
        if by_path["fingerprint"] != deck_fingerprint:
            # Rule B: file was edited
            library.update_deck_fingerprint(deck_id, deck_fingerprint)
        # Rule A (also covers B after the update)
        return deck_id

    by_fp = library.get_deck_by_fingerprint(deck_fingerprint)
    if by_fp is not None:
        deck_id = by_fp["deck_id"]
        # Rule C: file was moved/renamed
        library.update_deck_path(deck_id, path_str)
        return deck_id

    # Rule D: brand-new deck
    deck_id = str(uuid.uuid4())
    library.create_deck(deck_id, deck_fingerprint, path_str)
    return deck_id


# ---------------------------------------------------------------------------
# Internal helpers
# ---------------------------------------------------------------------------


def _compute_deck_fingerprint(pptx_path: Path) -> str:
    """Backward-compat shim — delegates to identity module."""
    return compute_deck_fingerprint(pptx_path)


def _extract_structural_facts(slide) -> dict:  # type: ignore[type-arg]
    """Extract text and shape counts from a python-pptx slide object."""
    text_objects = 0
    images = 0
    text_parts: list[str] = []

    for shape in slide.shapes:
        if shape.has_text_frame:
            text_objects += 1
            for para in shape.text_frame.paragraphs:
                t = para.text.strip()
                if t:
                    text_parts.append(t)

        if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
            images += 1

    charts = sum(1 for s in slide.shapes if getattr(s, "has_chart", False))
    tables = sum(1 for s in slide.shapes if getattr(s, "has_table", False))

    return {
        "text_object_count": text_objects,
        "image_count": images,
        "chart_count": charts,
        "table_count": tables,
        "shape_count": len(slide.shapes),
        "extracted_text": "\n".join(text_parts),
    }
