"""Stable slide identity: native PowerPoint IDs, UUID5 slide IDs, and fingerprints.

Functions in this module are pure — they read files and return values.
No Anthropic SDK, no SQLite, no pptx writing.
"""

from __future__ import annotations

import hashlib
import json
import posixpath
import uuid
import zipfile
import xml.etree.ElementTree as ET
from dataclasses import dataclass
from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

_NS_P = "http://schemas.openxmlformats.org/presentationml/2006/main"
_NS_R = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
_NS_REL = "http://schemas.openxmlformats.org/package/2006/relationships"
_SLIDE_REL_TYPE = (
    "http://schemas.openxmlformats.org/officeDocument/2006/relationships/slide"
)

# Relationship types whose referenced parts materially affect slide rendering.
# Theme, master, and layout references are excluded (deck-global, not slide-local).
_RENDERABLE_REL_TYPES = frozenset({
    "http://schemas.openxmlformats.org/officeDocument/2006/relationships/image",
    "http://schemas.openxmlformats.org/officeDocument/2006/relationships/chart",
    "http://schemas.openxmlformats.org/officeDocument/2006/relationships/diagramData",
    "http://schemas.openxmlformats.org/officeDocument/2006/relationships/diagramLayout",
    "http://schemas.openxmlformats.org/officeDocument/2006/relationships/diagramQuickStyle",
    "http://schemas.openxmlformats.org/officeDocument/2006/relationships/diagramColors",
})


@dataclass(frozen=True)
class SlideIdentity:
    """Native identity info for one slide in a PPTX deck."""
    slide_number: int       # 1-indexed, determined by order in sldIdLst
    native_slide_id: int    # from p:sldId id="" attribute — stable across edits
    slide_part: str         # zip-relative path: "slides/slideN.xml" (relative to ppt/)


def extract_slide_identities(pptx_path: Path) -> list[SlideIdentity]:
    """Return one SlideIdentity per slide, ordered by presentation order.

    Reads the PPTX zip directly (no python-pptx) so the native IDs are not
    altered.  The slide_part is relative to the ppt/ directory inside the zip,
    e.g. "slides/slide1.xml".

    Raises
    ------
    ValueError
        If the PPTX is malformed (missing presentation.xml or relationships).
    """
    with zipfile.ZipFile(pptx_path, "r") as z:
        # Parse relationship file to build {r_id -> slide_part} mapping.
        try:
            rels_xml = z.read("ppt/_rels/presentation.xml.rels")
        except KeyError as exc:
            raise ValueError(
                f"Malformed PPTX — missing ppt/_rels/presentation.xml.rels: {pptx_path}"
            ) from exc

        rels_root = ET.fromstring(rels_xml)
        r_id_to_part: dict[str, str] = {}
        for rel in rels_root.findall(f"{{{_NS_REL}}}Relationship"):
            if rel.get("Type") == _SLIDE_REL_TYPE:
                r_id = rel.get("Id", "")
                target = rel.get("Target", "")
                # Target is like "slides/slide1.xml" (relative to ppt/)
                if target.startswith("/ppt/"):
                    target = target[len("/ppt/"):]
                r_id_to_part[r_id] = target

        # Parse presentation.xml to get ordered sldIdLst entries.
        try:
            prs_xml = z.read("ppt/presentation.xml")
        except KeyError as exc:
            raise ValueError(
                f"Malformed PPTX — missing ppt/presentation.xml: {pptx_path}"
            ) from exc

        prs_root = ET.fromstring(prs_xml)
        sld_id_lst = prs_root.find(f"{{{_NS_P}}}sldIdLst")
        if sld_id_lst is None:
            return []

        identities: list[SlideIdentity] = []
        for slide_number, sld_id_elem in enumerate(
            sld_id_lst.findall(f"{{{_NS_P}}}sldId"), start=1
        ):
            native_id_str = sld_id_elem.get("id", "")
            r_id = sld_id_elem.get(f"{{{_NS_R}}}id", "")
            if not native_id_str or not r_id:
                continue
            slide_part = r_id_to_part.get(r_id)
            if slide_part is None:
                continue
            identities.append(
                SlideIdentity(
                    slide_number=slide_number,
                    native_slide_id=int(native_id_str),
                    slide_part=slide_part,
                )
            )

    return identities


def make_stable_slide_id(deck_id: str, native_slide_id: int) -> str:
    """Return a UUID5 string stable across text edits, reorders, and renames.

    UUID5 is deterministic: same (deck_id, native_slide_id) → same UUID.
    """
    return str(uuid.uuid5(uuid.NAMESPACE_URL, f"{deck_id}:{native_slide_id}"))


def compute_deck_fingerprint(pptx_path: Path) -> str:
    """SHA-256 of the raw file bytes; returned as a full 64-character hex string."""
    return hashlib.sha256(pptx_path.read_bytes()).hexdigest()


def _rels_zip_path(slide_part: str) -> str:
    """Return the ZIP path for the relationship sidecar of a slide part.

    "slides/slide1.xml"  →  "ppt/slides/_rels/slide1.xml.rels"
    """
    dir_part, file_part = (
        slide_part.rsplit("/", 1) if "/" in slide_part else ("", slide_part)
    )
    if dir_part:
        return f"ppt/{dir_part}/_rels/{file_part}.rels"
    return f"ppt/_rels/{file_part}.rels"


def _resolve_rel_target(slide_part: str, target: str) -> str:
    """Resolve a relationship target to an absolute ZIP path.

    slide_part : "slides/slide1.xml"
    target     : "../media/image1.png"  →  "ppt/media/image1.png"
    """
    if target.startswith("/"):
        return target.lstrip("/")
    slide_dir = "ppt/" + "/".join(slide_part.split("/")[:-1]) + "/"
    return posixpath.normpath(posixpath.join(slide_dir, target))


def compute_content_fingerprint(pptx_path: Path, slide_part: str) -> str:
    """SHA-256 fingerprint of a slide's visible content.

    Algorithm (Content Fingerprint v2):
        1.  slide_xml_hash  = SHA-256(raw bytes of ppt/<slide_part>)
        2.  Parse ppt/slides/_rels/<slideN>.xml.rels (if present).
        3.  For each Relationship whose Type is in _RENDERABLE_REL_TYPES
            (image, chart, diagramData, diagramLayout, diagramQuickStyle,
            diagramColors):
            a.  Resolve Target to an absolute ZIP path.
            b.  part_hash = SHA-256(referenced part bytes).
            c.  Collect entry (rel_type, resolved_zip_path, part_hash).
        4.  Sort entries lexicographically by (rel_type, resolved_zip_path).
        5.  If no entries:  return slide_xml_hash.
            (Unchanged from v1 for text-only slides — no backward-compat break.)
            Else: return SHA-256(
                slide_xml_hash
                + "::" + "::".join(f"{rt}:{p}:{h}" for rt, p, h in entries)
            )

    Theme, master, and layout references are excluded (deck-global assets whose
    changes should not invalidate individual slide classifications).
    """
    with zipfile.ZipFile(pptx_path, "r") as z:
        zip_names = set(z.namelist())

        # Step 1: hash raw slide XML.
        slide_bytes = z.read(f"ppt/{slide_part}")
        slide_xml_hash = hashlib.sha256(slide_bytes).hexdigest()

        # Step 2: parse slide relationship sidecar.
        rels_path = _rels_zip_path(slide_part)
        entries: list[tuple[str, str, str]] = []

        if rels_path in zip_names:
            rels_root = ET.fromstring(z.read(rels_path))
            for rel in rels_root.findall(f"{{{_NS_REL}}}Relationship"):
                rel_type = rel.get("Type", "")
                if rel_type not in _RENDERABLE_REL_TYPES:
                    continue
                target = rel.get("Target", "")
                resolved = _resolve_rel_target(slide_part, target)
                if resolved not in zip_names:
                    continue
                part_hash = hashlib.sha256(z.read(resolved)).hexdigest()
                entries.append((rel_type, resolved, part_hash))

        # Step 3: sort for determinism.
        entries.sort()

        # Step 4: combine.
        if not entries:
            return slide_xml_hash

        h = hashlib.sha256()
        h.update(slide_xml_hash.encode("utf-8"))
        for rel_type, resolved_path, part_hash in entries:
            h.update(f"::{rel_type}:{resolved_path}:{part_hash}".encode("utf-8"))
        return h.hexdigest()


def compute_structure_fingerprint(pptx_path: Path, slide_number: int) -> str:
    """SHA-256 of shape geometry — excludes all text values.

    The descriptor for each shape captures: shape_type, position, size,
    placeholder_type, has_chart, has_table.  Sorted by (top, left) for
    determinism.  Text changes do not affect this fingerprint.
    """
    prs = Presentation(str(pptx_path))
    slide = prs.slides[slide_number - 1]

    descriptors: list[dict] = []
    for shape in slide.shapes:
        placeholder_type: str | None = None
        if shape.is_placeholder:
            ph = shape.placeholder_format
            try:
                placeholder_type = ph.type.name
            except AttributeError:
                placeholder_type = str(ph.idx)

        try:
            shape_type_str = shape.shape_type.name
        except AttributeError:
            shape_type_str = str(shape.shape_type)

        descriptors.append(
            {
                "shape_type": shape_type_str,
                "left": int(shape.left) if shape.left is not None else None,
                "top": int(shape.top) if shape.top is not None else None,
                "width": int(shape.width) if shape.width is not None else None,
                "height": int(shape.height) if shape.height is not None else None,
                "placeholder_type": placeholder_type,
                "has_chart": bool(getattr(shape, "has_chart", False)),
                "has_table": bool(getattr(shape, "has_table", False)),
            }
        )

    # Sort by (top, left) so position-reordering in the XML does not matter.
    descriptors.sort(key=lambda d: (d["top"] or 0, d["left"] or 0))

    canonical = json.dumps(descriptors, sort_keys=True, separators=(",", ":"))
    return hashlib.sha256(canonical.encode("utf-8")).hexdigest()


def compute_input_fingerprint(
    slide_id: str,
    content_fp: str,
    structure_fp: str,
    classification_version: str,
    prompt_version: str,
) -> str:
    """SHA-256 of the colon-joined identity components.

    Changes whenever the slide content, structure, or classifier version changes.
    Used for staleness detection in the classification store.
    """
    raw = ":".join([slide_id, content_fp, structure_fp, classification_version, prompt_version])
    return hashlib.sha256(raw.encode("utf-8")).hexdigest()
