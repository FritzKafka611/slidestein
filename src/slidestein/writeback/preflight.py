"""Preflight validation for M6 write-back.

Performs ALL validation before any mutation:
  1. Completeness gate — draft.is_complete and no needs_input assignments.
  2. Draft / slot-map consistency — slide_id, deck_id, slide_number; per-slot identity.
  3. Output path safety — source != output; existing output vs overwrite flag.
  4. Source PPTX reachability — file exists, can be opened.
  5. Stable slide targeting — resolve current ordinal via deck_id + native slide ID (UUID5).
  6. Source-currentness check — recompute slot_analysis_input_fingerprint; reject if stale.
  7. Shape resolution — every shape_path resolves in live source at resolved ordinal.
  8. Shape ID validation — leaf shape_id matches.
  9. Before-write fingerprints — target structure + format fingerprints, deck-level baseline.

Returns a WritebackPlan if all checks pass.
Raises PowerPointWritebackError on the first failure found.
No mutation, no COM, no LLM.
"""

from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from slidestein.drafting.models import SlideContentDraft, SlotDraftAction
from slidestein.identity.slide_identity import (
    extract_slide_identities,
    make_stable_slide_id,
)
from slidestein.slots.fingerprint import (
    build_candidate_fingerprint_list,
    compute_slot_analysis_input_fingerprint,
)
from slidestein.slots.inspector import inspect_slide
from slidestein.slots.models import TemplateSlotMap
from slidestein.slots.versions import SLOT_ANALYSIS_PROMPT_VERSION, SLOT_MAP_SCHEMA_VERSION
from slidestein.writeback.errors import PowerPointWritebackError
from slidestein.writeback.models import (
    PowerPointWritebackRequest,
    WritebackOperation,
    WritebackPlan,
)
from slidestein.writeback.versions import WRITEBACK_SCHEMA_VERSION


# ---------------------------------------------------------------------------
# Shape-path resolution (python-pptx, read-only)
# ---------------------------------------------------------------------------


def _build_shape_path_map(
    shapes: object,
    parent_path: str | None = None,
) -> dict[str, object]:
    """Return {shape_path: pptx_shape} by recursively walking a ShapeCollection.

    shape_path format: "7" (top-level) or "12/7" or "31/46/48" (nested groups).
    Matches the convention in slots/inspector.py _walk_shapes().
    """
    result: dict[str, object] = {}
    for shape in shapes:  # type: ignore[union-attr]
        try:
            shape_id = int(shape.shape_id)
            path = f"{parent_path}/{shape_id}" if parent_path else str(shape_id)
            result[path] = shape
            try:
                if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
                    result.update(_build_shape_path_map(shape.shapes, path))
            except Exception:
                pass
        except Exception:
            continue
    return result


# ---------------------------------------------------------------------------
# Stable slide targeting
# ---------------------------------------------------------------------------


def resolve_slide_by_stable_id(
    source_pptx: Path,
    deck_id: str,
    target_slide_id: str,
) -> tuple[int, int]:
    """Return (resolved_slide_number, native_slide_id) for target_slide_id.

    Reads native slide IDs from the PPTX zip, computes the stable UUID5 for each,
    and returns the ordinal + native ID of the slide whose stable ID matches.

    Raises PowerPointWritebackError if no slide or multiple slides match.
    """
    identities = extract_slide_identities(source_pptx)
    matches = [
        si for si in identities
        if make_stable_slide_id(deck_id, si.native_slide_id) == target_slide_id
    ]

    if len(matches) == 0:
        slide_ids = [
            make_stable_slide_id(deck_id, si.native_slide_id)
            for si in identities
        ]
        raise PowerPointWritebackError(
            f"No slide in {source_pptx.name} resolves to "
            f"slot_map.slide_id={target_slide_id!r} "
            f"(deck_id={deck_id!r}, {len(identities)} slides checked: {slide_ids[:5]}). "
            "Source deck may have changed or wrong source provided."
        )

    if len(matches) > 1:
        raise PowerPointWritebackError(
            f"Multiple slides in {source_pptx.name} resolve to "
            f"slot_map.slide_id={target_slide_id!r} — ambiguous targeting. "
            f"Ordinals found: {[si.slide_number for si in matches]}"
        )

    return matches[0].slide_number, matches[0].native_slide_id


# ---------------------------------------------------------------------------
# Source-currentness check
# ---------------------------------------------------------------------------


def check_source_currentness(
    source_pptx: Path,
    resolved_slide_number: int,
    slot_map: TemplateSlotMap,
) -> None:
    """Fail before mutation if the source slide has changed since M5.2.

    Recomputes the slot_analysis_input_fingerprint from the current source PPTX
    and compares it to slot_map.slot_analysis_input_fingerprint.

    The fingerprint covers editable-shape geometry, text, and version constants —
    any change to shape positions, text, or addition/removal of editable shapes
    will produce a different fingerprint.

    Raises PowerPointWritebackError if the fingerprints do not match.
    """
    slide_id = slot_map.slide_id

    try:
        descriptors = inspect_slide(source_pptx, resolved_slide_number)
    except Exception as exc:
        raise PowerPointWritebackError(
            f"Cannot inspect slide {resolved_slide_number} of {source_pptx.name} "
            f"for source-currentness check: {exc}"
        ) from exc

    candidates = build_candidate_fingerprint_list(descriptors)

    current_fp = compute_slot_analysis_input_fingerprint(
        slide_id=slide_id,
        candidates=candidates,
        slot_map_schema_version=SLOT_MAP_SCHEMA_VERSION,
        prompt_version=SLOT_ANALYSIS_PROMPT_VERSION,
    )

    expected_fp = slot_map.slot_analysis_input_fingerprint

    if current_fp != expected_fp:
        raise PowerPointWritebackError(
            f"Source slide has changed since M5.2 analysis "
            f"(slide_id={slide_id!r}, slide={resolved_slide_number}). "
            f"slot_analysis_input_fingerprint mismatch: "
            f"expected={expected_fp[:16]}... current={current_fp[:16]}... "
            "Re-run M5.2 slot analysis or verify you supplied the correct source PPTX."
        )


# ---------------------------------------------------------------------------
# Public preflight function
# ---------------------------------------------------------------------------


def preflight(request: PowerPointWritebackRequest) -> WritebackPlan:
    """Validate the request and return a WritebackPlan ready for execution.

    All checks run before any COM or file mutation.

    Raises:
        PowerPointWritebackError: on any validation failure.
    """
    draft = request.draft
    slot_map = request.slot_map
    source_pptx = request.source_pptx
    output_pptx = request.output_pptx

    # ---- 1. Completeness gate ------------------------------------------------

    if not draft.is_complete:
        raise PowerPointWritebackError(
            f"Draft for slide '{draft.slide_id}' is not complete "
            f"(is_complete=False). M6 refuses incomplete drafts. "
            f"Complete all needs_input assignments before applying."
        )

    needs_input_slots = [
        a for a in draft.assignments if a.action == SlotDraftAction.NEEDS_INPUT
    ]
    if needs_input_slots:
        ids = [a.slot_id for a in needs_input_slots]
        raise PowerPointWritebackError(
            f"Draft for slide '{draft.slide_id}' contains {len(needs_input_slots)} "
            f"needs_input assignment(s) — cannot apply: {ids}"
        )

    # ---- 2. Draft / slot-map consistency -------------------------------------

    if draft.slide_id != slot_map.slide_id:
        raise PowerPointWritebackError(
            f"slide_id mismatch: draft.slide_id={draft.slide_id!r} "
            f"!= slot_map.slide_id={slot_map.slide_id!r}"
        )

    if draft.deck_id is not None and slot_map.deck_id is not None:
        if draft.deck_id != slot_map.deck_id:
            raise PowerPointWritebackError(
                f"deck_id mismatch: draft.deck_id={draft.deck_id!r} "
                f"!= slot_map.deck_id={slot_map.deck_id!r}"
            )

    if draft.slide_number != slot_map.slide_number:
        raise PowerPointWritebackError(
            f"slide_number mismatch: draft.slide_number={draft.slide_number} "
            f"!= slot_map.slide_number={slot_map.slide_number}"
        )

    # ---- 3. Per-assignment identity validation against slot map ---------------

    slot_by_id = {s.slot_id: s for s in slot_map.slots}

    # Check for duplicates in draft assignments
    seen_ids: set[str] = set()
    for a in draft.assignments:
        if a.slot_id in seen_ids:
            raise PowerPointWritebackError(
                f"Duplicate slot_id in draft assignments: {a.slot_id!r}"
            )
        seen_ids.add(a.slot_id)

    # Every assignment must match a slot in the slot map
    for assignment in draft.assignments:
        slot = slot_by_id.get(assignment.slot_id)
        if slot is None:
            raise PowerPointWritebackError(
                f"Assignment slot_id {assignment.slot_id!r} not found in slot map "
                f"(slide={draft.slide_id!r})"
            )
        if slot.shape_id != assignment.shape_id:
            raise PowerPointWritebackError(
                f"shape_id mismatch for slot {assignment.slot_id!r}: "
                f"assignment.shape_id={assignment.shape_id} "
                f"!= slot_map.shape_id={slot.shape_id} "
                f"(shape_path={slot.shape_path!r})"
            )
        if slot.shape_path != assignment.shape_path:
            raise PowerPointWritebackError(
                f"shape_path mismatch for slot {assignment.slot_id!r}: "
                f"assignment.shape_path={assignment.shape_path!r} "
                f"!= slot_map.shape_path={slot.shape_path!r}"
            )

    # Every slot map slot must have exactly one draft assignment
    draft_ids = {a.slot_id for a in draft.assignments}
    for slot in slot_map.slots:
        if slot.slot_id not in draft_ids:
            raise PowerPointWritebackError(
                f"Slot {slot.slot_id!r} (shape_path={slot.shape_path!r}) in slot map "
                f"has no corresponding draft assignment"
            )

    # ---- 4. Output path safety -----------------------------------------------

    try:
        src_abs = source_pptx.resolve()
        out_abs = output_pptx.resolve()
    except Exception:
        src_abs = source_pptx.absolute()
        out_abs = output_pptx.absolute()

    if src_abs == out_abs:
        raise PowerPointWritebackError(
            f"source_pptx and output_pptx resolve to the same file: {src_abs}. "
            "M6 will not overwrite the source."
        )

    if not request.overwrite and out_abs.exists():
        raise PowerPointWritebackError(
            f"Output file already exists: {out_abs}. "
            "Pass overwrite=True or --overwrite to replace it."
        )

    # ---- 5. Source PPTX reachability -----------------------------------------

    if not source_pptx.exists():
        raise PowerPointWritebackError(
            f"Source PPTX not found: {source_pptx}"
        )

    try:
        prs = Presentation(str(source_pptx))
    except Exception as exc:
        raise PowerPointWritebackError(
            f"Cannot open source PPTX {source_pptx}: {exc}"
        ) from exc

    # ---- 5b. Stable slide targeting ------------------------------------------
    # Requires deck_id to compute UUID5 stable slide ID.  Ordinal-only targeting
    # is not permitted per the M6 spec.

    if not slot_map.deck_id:
        raise PowerPointWritebackError(
            f"slot_map.deck_id is not set for slide {slot_map.slide_id!r}. "
            "deck_id is required for stable slide targeting. "
            "Ordinal-only targeting is not permitted."
        )

    try:
        resolved_slide_number, native_slide_id = resolve_slide_by_stable_id(
            source_pptx, slot_map.deck_id, slot_map.slide_id
        )
    except PowerPointWritebackError:
        raise
    except Exception as exc:
        raise PowerPointWritebackError(
            f"Stable slide targeting failed for {source_pptx.name}: {exc}"
        ) from exc

    # Validate resolved ordinal is in range (should always hold after resolve)
    if resolved_slide_number < 1 or resolved_slide_number > len(prs.slides):
        raise PowerPointWritebackError(
            f"Resolved slide ordinal {resolved_slide_number} out of range: "
            f"{source_pptx.name} has {len(prs.slides)} slide(s)"
        )

    pptx_slide = prs.slides[resolved_slide_number - 1]

    # ---- 6. Source-currentness check -----------------------------------------

    check_source_currentness(source_pptx, resolved_slide_number, slot_map)

    # ---- 7. Shape resolution against live source -----------------------------

    live_path_map = _build_shape_path_map(pptx_slide.shapes)

    for assignment in draft.assignments:
        shape_path = assignment.shape_path
        if shape_path not in live_path_map:
            raise PowerPointWritebackError(
                f"Cannot resolve shape_path {shape_path!r} for slot "
                f"{assignment.slot_id!r} on slide {resolved_slide_number} of "
                f"{source_pptx.name}"
            )
        live_shape = live_path_map[shape_path]
        live_id = int(live_shape.shape_id)
        if live_id != assignment.shape_id:
            raise PowerPointWritebackError(
                f"Shape at path {shape_path!r} has shape_id={live_id} in source, "
                f"expected shape_id={assignment.shape_id} "
                f"(slot {assignment.slot_id!r})"
            )
        if not getattr(live_shape, "has_text_frame", False):
            raise PowerPointWritebackError(
                f"Shape at path {shape_path!r} (shape_id={assignment.shape_id}) "
                f"does not have a text frame — cannot apply action "
                f"{assignment.action.value!r} for slot {assignment.slot_id!r}"
            )

    # ---- 8. Before-write fingerprints + deck baseline ------------------------

    from slidestein.identity.slide_identity import compute_structure_fingerprint
    from slidestein.writeback.verification import (
        compute_deck_structure_snapshot,
        compute_format_fingerprint,
    )

    try:
        before_fp = compute_structure_fingerprint(source_pptx, resolved_slide_number)
    except Exception:
        before_fp = ""

    try:
        before_fmt_fp = compute_format_fingerprint(source_pptx, resolved_slide_number)
    except Exception:
        before_fmt_fp = ""

    try:
        before_deck_native_ids, before_deck_structure = compute_deck_structure_snapshot(
            source_pptx
        )
    except Exception:
        before_deck_native_ids = []
        before_deck_structure = {}

    # ---- 9. Build WritebackPlan ----------------------------------------------

    operations: list[WritebackOperation] = []
    for assignment in draft.assignments:
        operations.append(WritebackOperation(
            slot_id=assignment.slot_id,
            shape_id=assignment.shape_id,
            shape_path=assignment.shape_path,
            action=assignment.action,
            text=assignment.text,
        ))

    replacement_count = sum(
        1 for op in operations if op.action == SlotDraftAction.REPLACE
    )
    clear_count = sum(
        1 for op in operations if op.action == SlotDraftAction.CLEAR
    )

    return WritebackPlan(
        schema_version=WRITEBACK_SCHEMA_VERSION,
        slide_id=draft.slide_id,
        deck_id=draft.deck_id,
        slide_number=slot_map.slide_number,
        resolved_slide_number=resolved_slide_number,
        native_slide_id=native_slide_id,
        source_pptx=source_pptx,
        output_pptx=output_pptx,
        operations=operations,
        replacement_count=replacement_count,
        clear_count=clear_count,
        before_structure_fingerprint=before_fp,
        before_format_fingerprint=before_fmt_fp,
        before_deck_native_ids=before_deck_native_ids,
        before_deck_structure=before_deck_structure,
    )
