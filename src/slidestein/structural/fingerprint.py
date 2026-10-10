"""Structural fingerprint for M10 — slot-layout identity.

A structural fingerprint commits to the semantic slot roles, normalized
slide-space geometry, and editable flags.  It deliberately excludes content
text so two slides with different text but identical layout are fingerprint-equal.

Two functions are provided:

compute_structural_fingerprint(slot_key_map)
    Uses TemplateSlot.geometry (x_ratio, y_ratio, etc.) as stored in the
    slot map.  For top-level shapes this equals slide-space.  For group-nested
    shapes this is the group-LOCAL coordinate divided by slide dimensions —
    NOT slide-space.  Use this only when no PPTX is available.

compute_structural_fingerprint_slide_space(pptx_path, slide_number, slot_key_map)
    Reads the actual slide-space EMU geometry by composing group transforms
    (via diagnosis._slide_space_geometry).  Normalises to ratios against slide
    dimensions so the result is comparable across slides of the same deck
    format.  Parent group movement changes these ratios, so two templates with
    identical group-local structure but different parent positions produce
    DIFFERENT fingerprints.  Falls back to compute_structural_fingerprint on
    any failure (bad PPTX, missing shape, etc.).

Used to detect exact-layout clones (noop exclusion in reselection) and as a
stable candidate identity in the result model.
"""

from __future__ import annotations

import hashlib
import json
from pathlib import Path
from typing import TYPE_CHECKING

if TYPE_CHECKING:
    from slidestein.slots.models import TemplateSlot, TemplateSlotMap

_FLOAT_PRECISION = 6


def _round(v: float) -> float:
    return round(v, _FLOAT_PRECISION)


def compute_structural_fingerprint(slot_key_map: dict[str, "TemplateSlot"]) -> str:
    """Return a SHA-256 fingerprint committing to slot role, geometry, and editable.

    Uses TemplateSlot.geometry (x_ratio, y_ratio, width_ratio, height_ratio).
    For top-level shapes these are slide-space ratios.  For group-nested shapes
    they are GROUP-LOCAL coordinates / slide dimensions — not slide-space.
    Use compute_structural_fingerprint_slide_space when authoritative
    slide-space geometry is required.

    Inputs (for each slot in K-key order):
      - slot_role.value
      - x_ratio, y_ratio, width_ratio, height_ratio  (6 d.p.)
      - editable (bool)

    Content text is intentionally excluded.
    """
    entries = []
    for key in sorted(slot_key_map.keys(), key=lambda k: int(k[1:])):
        slot = slot_key_map[key]
        geo = slot.geometry
        entries.append({
            "key": key,
            "role": slot.slot_role.value,
            "x_ratio": _round(geo.get("x_ratio", 0.0)),
            "y_ratio": _round(geo.get("y_ratio", 0.0)),
            "width_ratio": _round(geo.get("width_ratio", 0.0)),
            "height_ratio": _round(geo.get("height_ratio", 0.0)),
            "editable": bool(slot.editable),
        })

    payload = json.dumps(entries, sort_keys=True, ensure_ascii=False)
    return hashlib.sha256(payload.encode("utf-8")).hexdigest()


def compute_structural_fingerprint_from_slot_map(slot_map: "TemplateSlotMap") -> str:
    """Convenience wrapper accepting a TemplateSlotMap."""
    from slidestein.drafting.service import build_slot_key_mapping

    slot_key_map = build_slot_key_mapping(slot_map.slots)
    return compute_structural_fingerprint(slot_key_map)


def compute_structural_fingerprint_slide_space(
    pptx_path: Path,
    slide_number: int,
    slot_key_map: dict[str, "TemplateSlot"],
) -> str:
    """Return SHA-256 fingerprint using authoritative slide-space EMU geometry.

    Reads actual slide-space coordinates by composing group transforms at every
    ancestor level (via _slide_space_geometry).  Normalises to ratios against
    slide dimensions so the result is comparable across slides.

    Parent group movement changes the slide-space ratios, so two templates with
    identical group-local structure but different parent positions will produce
    DIFFERENT fingerprints.

    Falls back to compute_structural_fingerprint(slot_key_map) if:
      - PPTX cannot be opened or parsed
      - slide_number is out of range
      - any slot shape is missing from the path map

    No API or provider calls are made.
    """
    try:
        from pptx import Presentation  # noqa: PLC0415

        from slidestein.structural.diagnosis import (  # noqa: PLC0415
            _build_path_map,
            _slide_space_geometry,
        )

        prs = Presentation(str(pptx_path))
        n_slides = len(prs.slides)
        if slide_number < 1 or slide_number > n_slides:
            return compute_structural_fingerprint(slot_key_map)

        slide = prs.slides[slide_number - 1]
        slide_width = int(prs.slide_width)
        slide_height = int(prs.slide_height)

        if slide_width == 0 or slide_height == 0:
            return compute_structural_fingerprint(slot_key_map)

        path_map = _build_path_map(slide.shapes)

        entries = []
        for key in sorted(slot_key_map.keys(), key=lambda k: int(k[1:])):
            slot = slot_key_map[key]
            geo = _slide_space_geometry(path_map, slot.shape_path)
            if geo is None:
                # Shape missing or unresolvable — fall back to group-local version
                return compute_structural_fingerprint(slot_key_map)
            entries.append({
                "key": key,
                "role": slot.slot_role.value,
                "x_ratio": _round(geo["x"] / slide_width),
                "y_ratio": _round(geo["y"] / slide_height),
                "width_ratio": _round(geo["width"] / slide_width),
                "height_ratio": _round(geo["height"] / slide_height),
                "editable": bool(slot.editable),
            })

        payload = json.dumps(entries, sort_keys=True, ensure_ascii=False)
        return hashlib.sha256(payload.encode("utf-8")).hexdigest()

    except Exception:
        return compute_structural_fingerprint(slot_key_map)
