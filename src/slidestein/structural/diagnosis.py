"""StructuralDiagnosisService — deterministic structural diagnosis (zero API calls).

Compares original selected template slide geometry against the current
generated slide geometry for every semantic TemplateSlot.

Routing policy:
  - content identity mismatch           → StructuralRecoveryError (hard stop)
  - any semantic shape missing/drifted   → rebuild_current_template
  - no drift AND M9 escalated            → reselect_template
  - geometry resolution failure          → manual_review

Zero provider calls.  All inputs are validated before this service is called.

Geometry semantics
------------------
All EMU coordinates are in SLIDE SPACE.  For shapes nested inside one or more
group shapes the slide-space position is computed by composing the group
transforms at every ancestor level.

For a group G with:
  • position in parent space: (G.left, G.top)
  • child-coordinate origin:  (chOff_x, chOff_y)
  • parent-space extent:      (G.width, G.height)
  • child-space extent:       (chExt_cx, chExt_cy)

The transform from the child's LOCAL coordinates (cx, cy) to the PARENT's
LOCAL coordinates is:

  parent_x = G.left + (cx - chOff_x) * G.width  / chExt_cx
  parent_y = G.top  + (cy - chOff_y) * G.height / chExt_cy

When chExt equals the group extent (scale = 1.0, the typical case):

  parent_x = G.left + cx - chOff_x

Transforms compose for arbitrarily nested groups.
Width/height also scale proportionally.
"""

from __future__ import annotations

from pathlib import Path
from typing import TYPE_CHECKING, Optional

from slidestein.drafting.service import build_slot_key_mapping
from slidestein.structural.errors import StructuralRecoveryError
from slidestein.structural.models import (
    StructuralDrift,
    StructuralRecoveryPlan,
    StructuralRecoveryRoute,
)
from slidestein.writeback.errors import PowerPointWritebackError
from slidestein.writeback.preflight import resolve_slide_by_stable_id

if TYPE_CHECKING:
    from slidestein.slots.models import TemplateSlot, TemplateSlotMap
    from slidestein.structural.models import StructuralRecoveryRequest


# ---------------------------------------------------------------------------
# Group-transform helpers
# ---------------------------------------------------------------------------

_OOXML_DRAW_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"


def _get_group_xfrm(group_shape: object) -> tuple[int, int, int, int]:
    """Return (chOff_x, chOff_y, chExt_cx, chExt_cy) from group XML.

    Falls back to (0, 0, group.width, group.height) on any error so that
    the calling code degrades gracefully to a scale=1.0 / no-chOff transform.
    """
    try:
        elem = group_shape._element  # type: ignore[union-attr]
        ns = _OOXML_DRAW_NS
        xfrm = elem.find(f".//{{{ns}}}xfrm")
        if xfrm is None:
            return (0, 0, int(group_shape.width), int(group_shape.height))  # type: ignore[union-attr]

        chOff = xfrm.find(f"{{{ns}}}chOff")
        chExt = xfrm.find(f"{{{ns}}}chExt")

        off_x = int(chOff.get("x", 0)) if chOff is not None else 0
        off_y = int(chOff.get("y", 0)) if chOff is not None else 0
        ext_cx = int(chExt.get("cx") or group_shape.width) if chExt is not None else int(group_shape.width)  # type: ignore[union-attr]
        ext_cy = int(chExt.get("cy") or group_shape.height) if chExt is not None else int(group_shape.height)  # type: ignore[union-attr]
        return (off_x, off_y, ext_cx, ext_cy)
    except Exception:
        try:
            return (0, 0, int(group_shape.width), int(group_shape.height))  # type: ignore[union-attr]
        except Exception:
            return (0, 0, 1, 1)  # non-zero to avoid division by zero


# ---------------------------------------------------------------------------
# Shape path map builder
# ---------------------------------------------------------------------------


def _build_path_map(shapes: object, parent_path: Optional[str] = None) -> dict:
    """Recursively map shape_path → shape for all shapes on a slide.

    shape_path format: "21" (top-level) or "31/16" (group/child) etc.
    """
    from pptx.enum.shapes import MSO_SHAPE_TYPE  # noqa: PLC0415

    result: dict = {}
    try:
        for shape in shapes:  # type: ignore[union-attr]
            try:
                sid = int(shape.shape_id)
                path = f"{parent_path}/{sid}" if parent_path else str(sid)
                result[path] = shape
                try:
                    if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
                        result.update(_build_path_map(shape.shapes, path))
                except Exception:
                    pass
            except Exception:
                continue
    except Exception:
        pass
    return result


# ---------------------------------------------------------------------------
# Slide-space geometry computation
# ---------------------------------------------------------------------------


def _slide_space_geometry(path_map: dict, shape_path: str) -> Optional[dict]:
    """Compute slide-space EMU geometry by composing group transforms.

    Parameters
    ----------
    path_map:   mapping from shape_path to shape object (from _build_path_map)
    shape_path: slash-delimited path, e.g. "21" or "31/16" or "31/38/36"

    Returns
    -------
    {"x": int, "y": int, "width": int, "height": int} in slide EMU, or None
    if any ancestor or the target shape is missing from path_map.

    Algorithm
    ---------
    Maintain a cumulative affine transform (offset_x, offset_y, scale_x,
    scale_y) from slide origin.  At each intermediate GROUP level, compose
    that group's parent→child transform into the running values.  At the
    LEAF level, apply the final transform to the shape's local coordinates.

    For a group G:
      new_offset_x = offset_x + scale_x * (G.left - chOff_x * G.width / chExt_cx)
      new_scale_x  = scale_x  * G.width / chExt_cx

    (Similarly for y.)  When scale=1 and chOff=(0,0) this reduces to the
    simple addition: new_offset_x = offset_x + G.left.
    """
    parts = shape_path.split("/")
    offset_x: float = 0.0
    offset_y: float = 0.0
    scale_x: float = 1.0
    scale_y: float = 1.0

    for i, _part in enumerate(parts):
        current_path = "/".join(parts[: i + 1])
        shape = path_map.get(current_path)
        if shape is None:
            return None  # shape missing from slide

        if i < len(parts) - 1:
            # Intermediate group — compose its transform
            try:
                gx = int(shape.left)
                gy = int(shape.top)
                gw = int(shape.width)
                gh = int(shape.height)
            except Exception:
                return None

            ch_off_x, ch_off_y, ch_ext_cx, ch_ext_cy = _get_group_xfrm(shape)

            # Guard against zero extent (degenerate group)
            if ch_ext_cx == 0 or ch_ext_cy == 0:
                return None

            sx_factor = gw / ch_ext_cx
            sy_factor = gh / ch_ext_cy

            offset_x = offset_x + scale_x * (gx - ch_off_x * sx_factor)
            offset_y = offset_y + scale_y * (gy - ch_off_y * sy_factor)
            scale_x *= sx_factor
            scale_y *= sy_factor
        else:
            # Leaf shape — apply accumulated transform
            try:
                local_x = int(shape.left)
                local_y = int(shape.top)
                local_w = int(shape.width)
                local_h = int(shape.height)
            except Exception:
                return None

            return {
                "x": int(round(offset_x + scale_x * local_x)),
                "y": int(round(offset_y + scale_y * local_y)),
                "width": int(round(scale_x * local_w)),
                "height": int(round(scale_y * local_h)),
            }

    return None  # empty path (should not happen)


# ---------------------------------------------------------------------------
# PPTX geometry reader
# ---------------------------------------------------------------------------


def _read_slide_space_geometries(
    pptx_path: Path,
    slide_number: int,
    slot_key_map: dict[str, "TemplateSlot"],
) -> Optional[dict[str, Optional[dict]]]:
    """Read slide-space EMU geometry for each slot in slot_key_map.

    Returns a dict {key: geo_dict_or_None} where:
      - geo_dict_or_None is None  if the shape is missing or geometry unreadable
      - geo_dict has keys x, y, width, height in slide EMU

    Returns None (outer) if the PPTX cannot be opened or slide is out of range.
    """
    try:
        from pptx import Presentation  # noqa: PLC0415
    except Exception:
        return None

    try:
        prs = Presentation(str(pptx_path))
    except Exception:
        return None

    n_slides = len(prs.slides)
    if slide_number < 1 or slide_number > n_slides:
        return None

    pptx_slide = prs.slides[slide_number - 1]
    path_map = _build_path_map(pptx_slide.shapes)

    result: dict[str, Optional[dict]] = {}
    for key, slot in slot_key_map.items():
        geo = _slide_space_geometry(path_map, slot.shape_path)
        result[key] = geo  # None if missing/unreadable

    return result


# ---------------------------------------------------------------------------
# Drift classification
# ---------------------------------------------------------------------------


def _classify_drift(
    template_geo: dict,
    generated_geo: Optional[dict],
) -> Optional[str]:
    """Return drift_type string if geometries differ, None if equal.

    Inputs are slide-space EMU geometry dicts with keys x, y, width, height.
    Comparison is exact integer equality (no tolerance).
    """
    if not generated_geo:
        return "missing_shape"

    t_x, t_y = template_geo.get("x", 0), template_geo.get("y", 0)
    t_w, t_h = template_geo.get("width", 0), template_geo.get("height", 0)
    g_x, g_y = generated_geo.get("x", 0), generated_geo.get("y", 0)
    g_w, g_h = generated_geo.get("width", 0), generated_geo.get("height", 0)

    position_changed = (t_x != g_x) or (t_y != g_y)
    size_changed = (t_w != g_w) or (t_h != g_h)

    if position_changed and size_changed:
        return "position_and_size_changed"
    if position_changed:
        return "position_changed"
    if size_changed:
        return "size_changed"
    return None


# ---------------------------------------------------------------------------
# Service
# ---------------------------------------------------------------------------


class StructuralDiagnosisService:
    """Compares template and generated geometry deterministically.  Zero API calls."""

    def diagnose(self, request: "StructuralRecoveryRequest") -> StructuralRecoveryPlan:
        """Diagnose structural drift and determine the recovery route.

        Raises StructuralRecoveryError on:
          - generated content identity mismatch (hard stop — wrong evidence)

        Returns StructuralRecoveryPlan with route and any detected drifts.
        """
        slot_map = request.current_slot_map
        draft = request.current_draft

        slot_key_map = build_slot_key_mapping(slot_map.slots)

        # Resolve generated slide number
        if not slot_map.deck_id:
            return self._manual_review(
                request,
                ["slot_map.deck_id is not set — cannot resolve stable slide identity"],
            )

        try:
            generated_slide_number, _ = resolve_slide_by_stable_id(
                request.generated_pptx,
                slot_map.deck_id,
                slot_map.slide_id,
            )
        except PowerPointWritebackError:
            return self._manual_review(
                request,
                ["Cannot resolve generated slide: stable slide ID not found"],
            )

        # --- Generated content identity verification ---
        self._verify_content_identity(
            request.generated_pptx,
            generated_slide_number,
            slot_key_map,
            draft,
        )

        # --- Resolve template slide number ---
        try:
            template_slide_number, _ = resolve_slide_by_stable_id(
                request.template_pptx,
                slot_map.deck_id,
                slot_map.slide_id,
            )
        except PowerPointWritebackError:
            return self._manual_review(
                request,
                ["Cannot resolve template slide: stable slide ID not found"],
            )

        # --- Read template geometry (slide-space) ---
        template_geos = _read_slide_space_geometries(
            request.template_pptx, template_slide_number, slot_key_map
        )
        if template_geos is None:
            return self._manual_review(
                request,
                ["Cannot read template PPTX geometry"],
            )

        # --- Read generated geometry (slide-space) ---
        generated_geos = _read_slide_space_geometries(
            request.generated_pptx, generated_slide_number, slot_key_map
        )
        if generated_geos is None:
            return self._manual_review(
                request,
                ["Cannot read generated PPTX geometry"],
            )

        # --- Compare geometry per slot ---
        drifts: list[StructuralDrift] = []
        geometry_unresolvable = False

        for key, slot in slot_key_map.items():
            t_geo = template_geos.get(key)
            g_geo = generated_geos.get(key)

            # None from template means shape missing or unreadable
            if t_geo is None:
                geometry_unresolvable = True
                break

            drift_type = _classify_drift(t_geo, g_geo)
            if drift_type:
                drifts.append(
                    StructuralDrift(
                        key=key,
                        slot_id=slot.slot_id,
                        shape_path=slot.shape_path,
                        template_geometry=t_geo,
                        generated_geometry=g_geo,
                        drift_type=drift_type,
                    )
                )

        if geometry_unresolvable:
            return self._manual_review(
                request,
                ["Template geometry resolution failed for one or more semantic slots"],
            )

        return self._determine_route(request, drifts)

    def _verify_content_identity(
        self,
        generated_pptx: Path,
        slide_number: int,
        slot_key_map: dict,
        draft: "SlideContentDraft",  # type: ignore[name-defined]
    ) -> None:
        """Verify generated PPTX text matches the supplied draft.

        Raises StructuralRecoveryError if any assignment's text does not match.
        This is content provenance protection — M10 must not attempt recovery
        from stale or mismatched evidence.
        """
        from slidestein.drafting.models import SlotDraftAction  # noqa: PLC0415
        from slidestein.writeback.com_writer import normalize_readback  # noqa: PLC0415
        from slidestein.writeback.verification import read_shape_text  # noqa: PLC0415

        assignments_by_slot_id = {a.slot_id: a for a in draft.assignments}
        for key, slot in slot_key_map.items():
            assignment = assignments_by_slot_id[slot.slot_id]
            if assignment.action == SlotDraftAction.NEEDS_INPUT:
                continue

            actual_text = read_shape_text(generated_pptx, slide_number, slot.shape_path)

            if assignment.action == SlotDraftAction.REPLACE:
                expected = normalize_readback(assignment.text or "")
                if actual_text is None:
                    raise StructuralRecoveryError(
                        f"Content identity check failed for {key} "
                        f"(shape_path={slot.shape_path!r}): shape not found or "
                        "has no text frame. Generated PPTX does not correspond "
                        "to the supplied draft."
                    )
                if normalize_readback(actual_text) != expected:
                    raise StructuralRecoveryError(
                        f"Content identity check failed for {key} "
                        f"(slot_id={slot.slot_id!r}): generated text does not "
                        "match draft assignment. Generated PPTX does not "
                        "correspond to the supplied draft."
                    )
            elif assignment.action == SlotDraftAction.CLEAR:
                if actual_text is None:
                    raise StructuralRecoveryError(
                        f"Content identity check failed for {key} "
                        f"(shape_path={slot.shape_path!r}): shape not found. "
                        "A missing shape is never a successful CLEAR."
                    )
                if normalize_readback(actual_text) != "":
                    raise StructuralRecoveryError(
                        f"Content identity check failed for {key} "
                        f"(slot_id={slot.slot_id!r}): slot is marked CLEAR but "
                        "generated PPTX contains non-empty text."
                    )

    def _determine_route(
        self,
        request: "StructuralRecoveryRequest",
        drifts: list[StructuralDrift],
    ) -> StructuralRecoveryPlan:
        if drifts:
            drifted_keys = [d.key for d in drifts]
            reasons = [
                f"Generated artifact has structural drift from template in: "
                f"{', '.join(drifted_keys)}. "
                "Rebuild from original template is required."
            ]
            return StructuralRecoveryPlan(
                slide_id=request.current_slot_map.slide_id,
                deck_id=request.current_slot_map.deck_id,
                slide_number=request.current_slot_map.slide_number,
                route=StructuralRecoveryRoute.REBUILD_CURRENT_TEMPLATE,
                reasons=reasons,
                structural_drifts=drifts,
                current_template_slide_id=request.current_slot_map.slide_id,
                requires_candidate_retrieval=False,
                requires_selection_call=False,
                requires_redraft_call=False,
            )

        # No drift — template itself is the structural problem
        structural_issues = [
            iss for iss in request.visual_qa.issues
            if iss.category in (
                "visual_hierarchy",
                "alignment_and_spacing",
                "balance_and_whitespace",
                "typography_and_style_consistency",
            )
        ]
        reasons = [
            "Generated artifact geometry faithfully matches the selected template. "
            "M9 escalated to template_reselection due to material structural M8 issues. "
            "Template itself is a poor structural fit for this communication job."
        ]
        if structural_issues:
            cats = sorted({i.category for i in structural_issues})
            reasons.append(
                f"Material structural issues: {', '.join(cats)}."
            )
        return StructuralRecoveryPlan(
            slide_id=request.current_slot_map.slide_id,
            deck_id=request.current_slot_map.deck_id,
            slide_number=request.current_slot_map.slide_number,
            route=StructuralRecoveryRoute.RESELECT_TEMPLATE,
            reasons=reasons,
            structural_drifts=[],
            current_template_slide_id=request.current_slot_map.slide_id,
            requires_candidate_retrieval=True,
            requires_selection_call=True,
            requires_redraft_call=True,
        )

    @staticmethod
    def _manual_review(
        request: "StructuralRecoveryRequest",
        reasons: list[str],
    ) -> StructuralRecoveryPlan:
        return StructuralRecoveryPlan(
            slide_id=request.current_slot_map.slide_id,
            deck_id=request.current_slot_map.deck_id,
            slide_number=request.current_slot_map.slide_number,
            route=StructuralRecoveryRoute.MANUAL_REVIEW,
            reasons=reasons,
            structural_drifts=[],
            current_template_slide_id=request.current_slot_map.slide_id,
            requires_candidate_retrieval=False,
            requires_selection_call=False,
            requires_redraft_call=False,
        )
