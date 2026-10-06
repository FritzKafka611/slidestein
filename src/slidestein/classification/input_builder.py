"""Build a SlideClassificationInput from a PPTX file and slide number.

Slide ID strategy
-----------------
By default, slide_id = make_slide_id(deck_fingerprint, slide_number), which
matches the legacy identity format.  Pass slide_id explicitly to use the UUID5
identity assigned by the ingestion service (M3.4 stable identity).
"""

from __future__ import annotations

import hashlib
from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from slidestein.domain.models import SlideClassificationInput, make_slide_id


def build_slide_classification_input(
    pptx_path: Path,
    slide_number: int,
    preview_path: Path | None = None,
    slide_id: str | None = None,
) -> SlideClassificationInput:
    """Build a SlideClassificationInput for a single slide.

    Parameters
    ----------
    pptx_path:
        Path to the source PPTX.  Resolved to absolute before use.
    slide_number:
        1-indexed slide number.
    preview_path:
        Optional path to a pre-rendered PNG preview.  The file need not
        exist at call time; it is forwarded directly to the classifier.
    slide_id:
        Override the generated slide_id.  When None (default), the legacy
        ``make_slide_id(deck_fingerprint, slide_number)`` formula is used.
        Pass the UUID5 slide_id from a SlideRecord for persistence-aware flows.

    Raises
    ------
    FileNotFoundError
        If pptx_path does not exist.
    ValueError
        If pptx_path is not a .pptx file, cannot be opened, or
        slide_number is out of range.
    """
    pptx_path = pptx_path.resolve()

    if not pptx_path.exists():
        raise FileNotFoundError(f"PPTX file not found: {pptx_path}")
    if pptx_path.suffix.lower() != ".pptx":
        raise ValueError(f"Expected a .pptx file, got: {pptx_path.suffix!r}")

    try:
        prs = Presentation(str(pptx_path))
    except Exception as exc:
        raise ValueError(f"Cannot open PPTX file {pptx_path.name}: {exc}") from exc

    slide_count = len(prs.slides)
    if slide_number < 1 or slide_number > slide_count:
        raise ValueError(
            f"Slide {slide_number} does not exist; "
            f"presentation contains {slide_count} slide(s)."
        )

    if slide_id is None:
        deck_fingerprint = hashlib.sha256(pptx_path.read_bytes()).hexdigest()
        slide_id = make_slide_id(deck_fingerprint, slide_number)

    slide = prs.slides[slide_number - 1]
    extracted_text = _extract_text(slide)
    metadata = _build_structural_metadata(slide, slide_number, prs)

    return SlideClassificationInput(
        slide_id=slide_id,
        extracted_text=extracted_text,
        structural_metadata=metadata,
        preview_path=preview_path,
    )


# ---------------------------------------------------------------------------
# Internal helpers
# ---------------------------------------------------------------------------


def _extract_text(slide) -> str:  # type: ignore[type-arg]
    """Concatenate all non-empty paragraph texts from text-bearing shapes."""
    parts: list[str] = []
    for shape in slide.shapes:
        if shape.has_text_frame:
            for para in shape.text_frame.paragraphs:
                t = para.text.strip()
                if t:
                    parts.append(t)
    return "\n".join(parts)


def _build_structural_metadata(slide, slide_number: int, prs) -> dict:  # type: ignore[type-arg]
    """Build a JSON-serialisable structural metadata dict for one slide.

    Dimensions are in EMU (English Metric Units; 914400 EMU = 1 inch).
    All values are natively JSON-serialisable (int, str, list, dict, or None).
    """
    text_shapes: list[dict] = []
    text_shape_count = 0

    for shape in slide.shapes:
        if not shape.has_text_frame:
            continue
        text_shape_count += 1

        placeholder_type: str | None = None
        if shape.is_placeholder:
            ph = shape.placeholder_format
            try:
                placeholder_type = ph.type.name
            except AttributeError:
                placeholder_type = str(ph.idx)

        text_shapes.append(
            {
                "shape_id": shape.shape_id,
                "shape_name": shape.name,
                "text": shape.text_frame.text,
                "left": int(shape.left) if shape.left is not None else None,
                "top": int(shape.top) if shape.top is not None else None,
                "width": int(shape.width) if shape.width is not None else None,
                "height": int(shape.height) if shape.height is not None else None,
                "placeholder_type": placeholder_type,
            }
        )

    chart_count = sum(1 for s in slide.shapes if getattr(s, "has_chart", False))
    table_count = sum(1 for s in slide.shapes if getattr(s, "has_table", False))
    picture_count = sum(
        1 for s in slide.shapes if s.shape_type == MSO_SHAPE_TYPE.PICTURE
    )
    group_count = sum(
        1 for s in slide.shapes if s.shape_type == MSO_SHAPE_TYPE.GROUP
    )

    return {
        "slide_number": slide_number,
        "slide_width": int(prs.slide_width),
        "slide_height": int(prs.slide_height),
        "shape_count": len(slide.shapes),
        "text_shape_count": text_shape_count,
        "chart_count": chart_count,
        "table_count": table_count,
        "picture_count": picture_count,
        "group_count": group_count,
        "text_shapes": text_shapes,
    }
