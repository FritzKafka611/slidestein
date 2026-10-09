"""Slide shape inspector for M5.2 — deterministic, no LLM.

Walks all shapes on a slide via python-pptx (cross-platform, consistent with
the ingestion pipeline) and returns a list of NativeShapeDescriptors.

Groups are traversed recursively.  Each group child has its parent_shape_path
set to the parent's shape_path, and a shape_path of the form
"{parent_path}/{shape_id}".

Application-owned editability logic (can_edit_text) is determined here, not
by Vision.
"""

from __future__ import annotations

from pathlib import Path
from typing import Optional

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from slidestein.domain.models import SlideRecord
from slidestein.slots.models import NativeShapeDescriptor


_EMU_PER_PT: float = 12700.0
_DEFAULT_FONT_PT: float = 12.0


# ---------------------------------------------------------------------------
# Internal helpers
# ---------------------------------------------------------------------------


def _shape_type_name(shape: object) -> str:
    try:
        st = shape.shape_type  # type: ignore[attr-defined]
        name = getattr(st, "name", None)
        if name:
            return str(name)
        return f"TYPE_{int(st)}"
    except Exception:
        return "UNKNOWN"


def _extract_text(shape: object) -> str:
    try:
        if getattr(shape, "has_text_frame", False):
            return shape.text_frame.text or ""  # type: ignore[union-attr]
        return ""
    except Exception:
        return ""


def _get_font_size_pt(shape: object) -> Optional[float]:
    try:
        if not getattr(shape, "has_text_frame", False):
            return None
        for para in shape.text_frame.paragraphs:  # type: ignore[union-attr]
            for run in para.runs:
                sz = run.font.size
                if sz is not None:
                    return float(sz) / _EMU_PER_PT
            sz = para.font.size
            if sz is not None:
                return float(sz) / _EMU_PER_PT
        return None
    except Exception:
        return None


def _safe_int(value: object, default: int = 0) -> int:
    try:
        return int(value) if value is not None else default
    except (TypeError, ValueError):
        return default


def _classify_source_kind(shape: object) -> str:
    try:
        if getattr(shape, "shape_type", None) is not None:
            if shape.shape_type == MSO_SHAPE_TYPE.GROUP:  # type: ignore[attr-defined]
                return "group"
        if getattr(shape, "has_chart", False):
            return "chart"
        if getattr(shape, "has_table", False):
            return "table"
        try:
            if shape.shape_type in (  # type: ignore[attr-defined]
                MSO_SHAPE_TYPE.PICTURE,
                MSO_SHAPE_TYPE.LINKED_PICTURE,
            ):
                return "image"
        except Exception:
            pass
        if getattr(shape, "is_placeholder", False) and getattr(shape, "has_text_frame", False):
            return "placeholder"
        if getattr(shape, "has_text_frame", False):
            return "text_box"
        return "other"
    except Exception:
        return "other"


def _can_edit_text(shape: object, source_kind: str) -> bool:
    # Application logic only — Vision cannot override this determination.
    if source_kind in ("group", "chart", "table", "image", "other"):
        return False
    return bool(getattr(shape, "has_text_frame", False))


def _walk_shapes(
    shapes: object,
    parent_path: Optional[str],
    z_base: int,
    slide_width: int,
    slide_height: int,
) -> list[NativeShapeDescriptor]:
    """Recursively walk a ShapeCollection, returning descriptors."""
    results: list[NativeShapeDescriptor] = []
    try:
        shape_list = list(shapes)  # type: ignore[arg-type]
    except Exception:
        return results

    for z_idx, shape in enumerate(shape_list):
        try:
            shape_id = _safe_int(getattr(shape, "shape_id", 0))
            shape_path = (
                f"{parent_path}/{shape_id}" if parent_path else str(shape_id)
            )

            left = _safe_int(getattr(shape, "left", 0))
            top = _safe_int(getattr(shape, "top", 0))
            width = _safe_int(getattr(shape, "width", 0))
            height = _safe_int(getattr(shape, "height", 0))

            x_ratio = left / slide_width if slide_width else 0.0
            y_ratio = top / slide_height if slide_height else 0.0
            w_ratio = width / slide_width if slide_width else 0.0
            h_ratio = height / slide_height if slide_height else 0.0

            source_kind = _classify_source_kind(shape)
            is_group = source_kind == "group"
            has_tf = bool(getattr(shape, "has_text_frame", False))
            text = _extract_text(shape)
            is_ph = bool(getattr(shape, "is_placeholder", False))

            ph_type: Optional[int] = None
            ph_idx: Optional[int] = None
            if is_ph:
                try:
                    pf = shape.placeholder_format  # type: ignore[attr-defined]
                    ph_type = int(pf.type) if pf.type is not None else None
                    ph_idx = int(pf.idx)
                except Exception:
                    pass

            has_table = bool(getattr(shape, "has_table", False))
            table_rows: Optional[int] = None
            table_cols: Optional[int] = None
            if has_table:
                try:
                    tbl = shape.table  # type: ignore[attr-defined]
                    table_rows = len(tbl.rows)
                    table_cols = len(tbl.columns)
                except Exception:
                    pass

            font_pt = _get_font_size_pt(shape)

            desc = NativeShapeDescriptor(
                shape_id=shape_id,
                shape_path=shape_path,
                shape_type=_shape_type_name(shape),
                x=left,
                y=top,
                width=width,
                height=height,
                x_ratio=round(x_ratio, 6),
                y_ratio=round(y_ratio, 6),
                width_ratio=round(w_ratio, 6),
                height_ratio=round(h_ratio, 6),
                z_order=z_base + z_idx,
                has_text=bool(text),
                text=text,
                is_placeholder=is_ph,
                placeholder_type=ph_type,
                placeholder_idx=ph_idx,
                is_group=is_group,
                parent_shape_path=parent_path,
                is_table=has_table,
                table_rows=table_rows,
                table_columns=table_cols,
                has_text_frame=has_tf,
                font_size_pt=font_pt,
                can_edit_text=_can_edit_text(shape, source_kind),
                source_kind=source_kind,
            )
            results.append(desc)

            # Recurse into groups
            if is_group:
                try:
                    child_shapes = shape.shapes  # type: ignore[attr-defined]
                    child_z_base = (z_base + z_idx + 1) * 1000
                    results.extend(
                        _walk_shapes(
                            child_shapes,
                            shape_path,
                            child_z_base,
                            slide_width,
                            slide_height,
                        )
                    )
                except Exception:
                    pass  # unsupported group type — skip children silently

        except Exception:
            continue  # skip malformed shapes rather than crashing

    return results


# ---------------------------------------------------------------------------
# Public interface
# ---------------------------------------------------------------------------


class SlideShapeInspector:
    """Extract NativeShapeDescriptors from a SlideRecord.

    Uses python-pptx for cross-platform reading.  Does not modify any file.
    Does not call any LLM.
    """

    def inspect(self, record: SlideRecord) -> list[NativeShapeDescriptor]:
        """Walk all shapes on the slide and return one descriptor per shape.

        The slide is loaded from ``record.source_deck_path`` at
        ``record.slide_number`` (1-indexed).  Groups are traversed
        recursively; group children appear after their parent.

        Raises:
            FileNotFoundError: if the source deck file does not exist.
            ValueError: if the slide number is out of range.
            RuntimeError: if python-pptx fails to open the file.
        """
        source_path = Path(str(record.source_deck_path))
        if not source_path.exists():
            raise FileNotFoundError(
                f"Source deck not found: {source_path}"
            )

        try:
            prs = Presentation(str(source_path))
        except Exception as exc:
            raise RuntimeError(
                f"Failed to open {source_path}: {type(exc).__name__}: {exc}"
            ) from exc

        slide_number = record.slide_number  # 1-indexed
        if slide_number < 1 or slide_number > len(prs.slides):
            raise ValueError(
                f"Slide number {slide_number} out of range "
                f"(deck has {len(prs.slides)} slide(s))"
            )

        slide = prs.slides[slide_number - 1]
        slide_width = _safe_int(getattr(prs.slide_width, "__int__", lambda: 0)(), 9144000)
        slide_height = _safe_int(getattr(prs.slide_height, "__int__", lambda: 0)(), 5143500)

        # Better: access directly
        try:
            slide_width = int(prs.slide_width)
            slide_height = int(prs.slide_height)
        except Exception:
            slide_width = 9144000   # 10 inches default
            slide_height = 5143500  # 5.625 inches default

        return _walk_shapes(slide.shapes, None, 0, slide_width, slide_height)
