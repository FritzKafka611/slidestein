"""Pillow-based K-key slot-locator overlay for M8 Visual QA.

Creates a copy of the source image with bounding boxes and K-key labels
drawn over each slot's LIVE GENERATED geometry.  Used as IMAGE 3 in the
Vision request — locator reference only, not sent back to the model for
scoring in place of the clean render.

Live geometry is read from the actual generated PPTX (not TemplateSlot.geometry
which reflects historical M5.2/template measurement) so the overlay
accurately locates shapes even when M6 write-back altered their positions.

If live geometry cannot be read for any reason (unreadable PPTX, invalid slide
number, missing shape, unreadable geometry property) LiveGeometryError is raised.
TemplateSlot.geometry is NEVER used as a fallback — historical template geometry
must never masquerade as generated PPTX geometry.
"""

from __future__ import annotations

from pathlib import Path
from typing import TYPE_CHECKING, Optional

from slidestein.qa.errors import LiveGeometryError

if TYPE_CHECKING:
    from slidestein.slots.models import TemplateSlot


def _build_path_map(shapes: object, parent: Optional[str] = None) -> dict:
    """Recursively build {shape_path: shape} from a python-pptx shapes collection."""
    from pptx.enum.shapes import MSO_SHAPE_TYPE  # noqa: PLC0415

    result: dict = {}
    try:
        for shape in shapes:  # type: ignore[union-attr]
            try:
                sid = int(shape.shape_id)
                path = f"{parent}/{sid}" if parent else str(sid)
                result[path] = shape
                try:
                    if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
                        result.update(_build_path_map(shape.shapes, path))
                except Exception:
                    pass
            except Exception:
                continue
    except Exception:
        pass
    return result


def read_live_slot_geometries(
    pptx_path: Path,
    slide_number: int,
    slot_key_map: dict[str, "TemplateSlot"],
) -> dict[str, dict]:
    """Read actual shape geometry from the generated PPTX for each K-key slot.

    Returns {key: {x_ratio, y_ratio, width_ratio, height_ratio}} using the
    actual left/top/width/height from the generated PPTX normalised by slide
    dimensions.

    Raises LiveGeometryError for any of:
    - PPTX cannot be opened
    - slide_number out of range
    - slide dimensions invalid (<= 0)
    - a slot's shape_path is not found in the slide
    - shape geometry properties cannot be read

    TemplateSlot.geometry is NEVER used as a fallback.
    """
    from pptx import Presentation  # noqa: PLC0415

    try:
        prs = Presentation(str(pptx_path))
    except Exception as exc:
        raise LiveGeometryError(
            f"Cannot open generated PPTX {pptx_path.name!r}: {exc}"
        ) from exc

    n_slides = len(prs.slides)
    if slide_number < 1 or slide_number > n_slides:
        raise LiveGeometryError(
            f"Slide number {slide_number} is out of range for {pptx_path.name!r} "
            f"(file has {n_slides} slides)"
        )

    try:
        slide_w = int(prs.slide_width)
        slide_h = int(prs.slide_height)
    except Exception as exc:
        raise LiveGeometryError(
            f"Cannot read slide dimensions from {pptx_path.name!r}: {exc}"
        ) from exc

    if slide_w <= 0 or slide_h <= 0:
        raise LiveGeometryError(
            f"Invalid slide dimensions {slide_w}x{slide_h} in {pptx_path.name!r}"
        )

    pptx_slide = prs.slides[slide_number - 1]
    path_map = _build_path_map(pptx_slide.shapes)

    live_geo: dict[str, dict] = {}
    for key, slot in slot_key_map.items():
        shape = path_map.get(slot.shape_path)
        if shape is None:
            raise LiveGeometryError(
                f"Shape {slot.shape_path!r} (key={key}, slot_id={slot.slot_id!r}) "
                f"not found in slide {slide_number} of {pptx_path.name!r}. "
                "Generated PPTX does not contain the expected shape."
            )
        try:
            left = int(shape.left)
            top = int(shape.top)
            width = int(shape.width)
            height = int(shape.height)
        except Exception as exc:
            raise LiveGeometryError(
                f"Cannot read geometry for shape {slot.shape_path!r} "
                f"(key={key}, slot_id={slot.slot_id!r}): {exc}"
            ) from exc
        live_geo[key] = {
            "x_ratio": left / slide_w,
            "y_ratio": top / slide_h,
            "width_ratio": width / slide_w,
            "height_ratio": height / slide_h,
        }

    return live_geo


def create_slot_overlay(
    source_image: Path,
    output_path: Path,
    slot_geometries: dict[str, dict],
) -> None:
    """Draw K-key bounding boxes on a copy of source_image and save to output_path.

    slot_geometries maps K-key -> {x_ratio, y_ratio, width_ratio, height_ratio}
    — these must come from the LIVE GENERATED PPTX, not TemplateSlot.geometry.
    Image dimensions are read from the source image at runtime so no EMU
    conversion is needed.

    Does NOT modify the source image.
    """
    from PIL import Image, ImageDraw, ImageFont  # noqa: PLC0415

    img = Image.open(str(source_image)).convert("RGB")
    overlay = img.copy()
    draw = ImageDraw.Draw(overlay)

    img_w, img_h = overlay.size

    try:
        font = ImageFont.load_default(size=14)
    except TypeError:
        font = ImageFont.load_default()

    box_colour = (220, 30, 30)
    label_bg = (220, 30, 30)
    label_fg = (255, 255, 255)

    for key, geo in sorted(slot_geometries.items(), key=lambda kv: int(kv[0][1:])):
        x_ratio = float(geo.get("x_ratio", 0.0))
        y_ratio = float(geo.get("y_ratio", 0.0))
        w_ratio = float(geo.get("width_ratio", 0.0))
        h_ratio = float(geo.get("height_ratio", 0.0))

        x0 = int(x_ratio * img_w)
        y0 = int(y_ratio * img_h)
        x1 = int((x_ratio + w_ratio) * img_w)
        y1 = int((y_ratio + h_ratio) * img_h)

        # Clamp to image bounds
        x0 = max(0, min(x0, img_w - 1))
        y0 = max(0, min(y0, img_h - 1))
        x1 = max(x0 + 1, min(x1, img_w))
        y1 = max(y0 + 1, min(y1, img_h))

        draw.rectangle([x0, y0, x1, y1], outline=box_colour, width=2)

        # Label background pill
        label_text = key
        try:
            bbox = draw.textbbox((0, 0), label_text, font=font)
            tw = bbox[2] - bbox[0]
            th = bbox[3] - bbox[1]
        except AttributeError:
            tw, th = draw.textsize(label_text, font=font)  # type: ignore[attr-defined]

        lx0 = x0
        ly0 = max(0, y0 - th - 4)
        lx1 = lx0 + tw + 6
        ly1 = ly0 + th + 4

        draw.rectangle([lx0, ly0, lx1, ly1], fill=label_bg)
        draw.text((lx0 + 3, ly0 + 2), label_text, fill=label_fg, font=font)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    overlay.save(str(output_path), format="PNG")
