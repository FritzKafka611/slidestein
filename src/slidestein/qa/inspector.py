"""Native deterministic visual inspection for M8 Visual QA.

Inspects the generated PPTX using python-pptx (always available) and
optionally PowerPoint COM (Windows + Office) to produce structured
NativeVisualCheck facts before the Vision call.

Checks implemented
------------------
outside_slide_bounds  — shape geometry extends outside slide area (python-pptx)
text_overflow         — text content exceeds text-frame height (COM, fallback: unknown)

The inspector is intentionally conservative: it only emits FAIL when the
evidence is unambiguous.  Visual overlap between intentional decorative
elements is NOT checked here — that is Vision's domain.

COM shape resolution reuses slidestein.writeback.com_writer.resolve_com_shape_by_path
(the authoritative M6 implementation) rather than maintaining a second resolver.
"""

from __future__ import annotations

from pathlib import Path
from typing import TYPE_CHECKING, Optional

from slidestein.qa.models import NativeCheckStatus, NativeCheckType, NativeVisualCheck

if TYPE_CHECKING:
    from slidestein.slots.models import TemplateSlot, TemplateSlotMap


# ---------------------------------------------------------------------------
# Internal helpers — python-pptx
# ---------------------------------------------------------------------------


def _build_shape_path_map_pptx(shapes: object, parent_path: Optional[str] = None) -> dict:
    """Walk python-pptx ShapeCollection, returning {shape_path: shape}."""
    from pptx.enum.shapes import MSO_SHAPE_TYPE  # noqa: PLC0415

    result: dict = {}
    try:
        shape_list = list(shapes)  # type: ignore[arg-type]
    except Exception:
        return result

    for shape in shape_list:
        try:
            shape_id = int(shape.shape_id)
            path = f"{parent_path}/{shape_id}" if parent_path else str(shape_id)
            result[path] = shape
            try:
                if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
                    result.update(_build_shape_path_map_pptx(shape.shapes, path))
            except Exception:
                pass
        except Exception:
            continue

    return result


def _check_outside_bounds(
    shape: object,
    slide_width_emu: int,
    slide_height_emu: int,
    slot_key: str,
    shape_path: str,
) -> NativeVisualCheck:
    """Return outside_slide_bounds check for one shape (python-pptx).

    Details include explicit overflow magnitudes in EMU for each overflowing edge.
    """
    try:
        left = int(shape.left)  # type: ignore[union-attr]
        top = int(shape.top)  # type: ignore[union-attr]
        width = int(shape.width)  # type: ignore[union-attr]
        height = int(shape.height)  # type: ignore[union-attr]
    except Exception:
        return NativeVisualCheck(
            check_type=NativeCheckType.OUTSIDE_SLIDE_BOUNDS,
            slot_key=slot_key,
            status=NativeCheckStatus.UNKNOWN,
            details=f"Could not read geometry for shape {shape_path!r}",
        )

    right = left + width
    bottom = top + height

    overflow_parts = []
    if left < 0:
        overflow_parts.append(f"left={left} EMU (overflow={abs(left)} EMU)")
    if top < 0:
        overflow_parts.append(f"top={top} EMU (overflow={abs(top)} EMU)")
    if right > slide_width_emu:
        overflow_parts.append(
            f"right={right} > slide_width={slide_width_emu} EMU (overflow={right - slide_width_emu} EMU)"
        )
    if bottom > slide_height_emu:
        overflow_parts.append(
            f"bottom={bottom} > slide_height={slide_height_emu} EMU (overflow={bottom - slide_height_emu} EMU)"
        )

    if overflow_parts:
        return NativeVisualCheck(
            check_type=NativeCheckType.OUTSIDE_SLIDE_BOUNDS,
            slot_key=slot_key,
            status=NativeCheckStatus.FAIL,
            details=f"Shape {shape_path!r} extends outside slide: {'; '.join(overflow_parts)}",
        )

    return NativeVisualCheck(
        check_type=NativeCheckType.OUTSIDE_SLIDE_BOUNDS,
        slot_key=slot_key,
        status=NativeCheckStatus.PASS,
        details=(
            f"Shape {shape_path!r} within slide bounds "
            f"(left={left}, top={top}, right={right}, bottom={bottom}, "
            f"slide={slide_width_emu}x{slide_height_emu} EMU)"
        ),
    )


# ---------------------------------------------------------------------------
# COM-based text overflow check
# ---------------------------------------------------------------------------


def _com_text_overflow_for_shapes(
    pptx_path: Path,
    slide_number: int,
    shape_paths: list[str],
) -> dict[str, Optional[bool]]:
    """Use PowerPoint COM to check text overflow for a list of shape paths.

    Returns dict mapping shape_path -> True (overflow), False (no overflow),
    or None (unknown / not determinable).

    Opens PowerPoint once, checks all shapes, closes.  Any exception on the
    outer level returns all shapes as None.

    COM overflow invariant (confirmed on this machine):
        TextFrame.Height does NOT exist in this COM version.
        Use Shape.Height as the authoritative height dimension.

        available_height = Shape.Height - tf.MarginTop - tf.MarginBottom
        overflow = tf.TextRange.BoundHeight > available_height

        When tf.AutoSize == 1 (shape auto-resizes to fit text):
            overflow is structurally impossible; returns False.

    Shape resolution reuses resolve_com_shape_by_path from
    slidestein.writeback.com_writer — the authoritative M6 implementation.
    """
    result: dict[str, Optional[bool]] = {p: None for p in shape_paths}
    try:
        import comtypes.client  # noqa: PLC0415

        from slidestein.writeback.com_writer import (  # noqa: PLC0415
            resolve_com_shape_by_path,
        )

        ppt = comtypes.client.CreateObject("PowerPoint.Application")
        ppt.Visible = 1
        try:
            prs = ppt.Presentations.Open(
                str(pptx_path.resolve()),
                ReadOnly=True,
                WithWindow=False,
            )
            try:
                slide = prs.Slides(slide_number)
                for sp in shape_paths:
                    try:
                        com_shape = resolve_com_shape_by_path(slide, sp)
                        if com_shape is None or not com_shape.HasTextFrame:
                            continue
                        tf = com_shape.TextFrame
                        # ppAutoSizeShapeToFitText = 1 — shape resizes; no overflow
                        if tf.AutoSize == 1:
                            result[sp] = False
                            continue
                        # Use Shape.Height (TextFrame.Height does not exist in this COM API)
                        available_pt = com_shape.Height - tf.MarginTop - tf.MarginBottom
                        actual_pt = tf.TextRange.BoundHeight
                        result[sp] = actual_pt > available_pt
                    except Exception:
                        result[sp] = None
            finally:
                prs.Close()
        finally:
            ppt.Quit()
    except Exception:
        pass  # result dict stays all None

    return result


# ---------------------------------------------------------------------------
# Public inspector
# ---------------------------------------------------------------------------


class NativeVisualInspector:
    """Deterministic native inspection of the generated PPTX for M8.

    Uses python-pptx for geometry / bounds checks (always available) and
    PowerPoint COM for text-overflow checks (Windows + Office; falls back to
    UNKNOWN automatically).
    """

    def inspect(
        self,
        pptx_path: Path,
        slide_number: int,
        slot_key_map: dict[str, "TemplateSlot"],
        slot_map: "TemplateSlotMap",
    ) -> list[NativeVisualCheck]:
        """Return all NativeVisualChecks for the targeted semantic slots."""
        from pptx import Presentation  # noqa: PLC0415

        checks: list[NativeVisualCheck] = []

        if not pptx_path.exists():
            checks.append(NativeVisualCheck(
                check_type=NativeCheckType.RENDERING_IDENTITY,
                slot_key=None,
                status=NativeCheckStatus.FAIL,
                details=f"Generated PPTX not found: {pptx_path}",
            ))
            return checks

        try:
            prs = Presentation(str(pptx_path))
        except Exception as exc:
            checks.append(NativeVisualCheck(
                check_type=NativeCheckType.RENDERING_IDENTITY,
                slot_key=None,
                status=NativeCheckStatus.FAIL,
                details=f"Cannot open generated PPTX: {type(exc).__name__}",
            ))
            return checks

        if slide_number < 1 or slide_number > len(prs.slides):
            checks.append(NativeVisualCheck(
                check_type=NativeCheckType.RENDERING_IDENTITY,
                slot_key=None,
                status=NativeCheckStatus.FAIL,
                details=(
                    f"Slide number {slide_number} out of range "
                    f"(deck has {len(prs.slides)} slides)"
                ),
            ))
            return checks

        try:
            slide_w = int(prs.slide_width)
            slide_h = int(prs.slide_height)
        except Exception:
            slide_w = slot_map.slide_width
            slide_h = slot_map.slide_height

        pptx_slide = prs.slides[slide_number - 1]
        shape_path_map = _build_shape_path_map_pptx(pptx_slide.shapes)

        # Build: shape_path -> slot_key
        path_to_key: dict[str, str] = {
            slot.shape_path: key for key, slot in slot_key_map.items()
        }

        # Paths with text frames for COM overflow check
        text_frame_paths: list[str] = []

        for shape_path, slot_key in sorted(path_to_key.items()):
            pptx_shape = shape_path_map.get(shape_path)
            if pptx_shape is None:
                checks.append(NativeVisualCheck(
                    check_type=NativeCheckType.OUTSIDE_SLIDE_BOUNDS,
                    slot_key=slot_key,
                    status=NativeCheckStatus.UNKNOWN,
                    details=f"Shape path {shape_path!r} not found in generated PPTX",
                ))
                continue

            # Bounds check (python-pptx, reliable)
            checks.append(_check_outside_bounds(pptx_shape, slide_w, slide_h, slot_key, shape_path))

            # Collect for text overflow (COM)
            if getattr(pptx_shape, "has_text_frame", False):
                text_frame_paths.append(shape_path)

        # Text overflow via COM (single PowerPoint session for all shapes)
        if text_frame_paths:
            overflow_results = _com_text_overflow_for_shapes(
                pptx_path, slide_number, text_frame_paths
            )
            for shape_path, overflows in overflow_results.items():
                slot_key = path_to_key[shape_path]
                if overflows is None:
                    checks.append(NativeVisualCheck(
                        check_type=NativeCheckType.TEXT_OVERFLOW,
                        slot_key=slot_key,
                        status=NativeCheckStatus.UNKNOWN,
                        details=(
                            f"Text overflow not determinable for {shape_path!r} "
                            "(COM inspection unavailable or failed)"
                        ),
                    ))
                elif overflows:
                    checks.append(NativeVisualCheck(
                        check_type=NativeCheckType.TEXT_OVERFLOW,
                        slot_key=slot_key,
                        status=NativeCheckStatus.FAIL,
                        details=(
                            f"Text content height exceeds available text-frame height "
                            f"for shape {shape_path!r} "
                            "(BoundHeight > Shape.Height - MarginTop - MarginBottom)"
                        ),
                    ))
                else:
                    checks.append(NativeVisualCheck(
                        check_type=NativeCheckType.TEXT_OVERFLOW,
                        slot_key=slot_key,
                        status=NativeCheckStatus.PASS,
                        details=f"Text fits within text-frame for shape {shape_path!r}",
                    ))

        return checks
