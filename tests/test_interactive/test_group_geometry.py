"""Tests for group geometry detection via slide-space coordinates."""

from __future__ import annotations

from slidestein.interactive.adapters.python_pptx import PythonPptxPowerPointHostAdapter
from slidestein.interactive.diff import SlideStateDiffService
from slidestein.structural.diagnosis import _build_path_map
from tests.test_interactive.conftest import (
    _GROUP_ID,
    _GROUP_X,
    make_pingpong_pptx,
    simulate_human_group_move,
)

_ADAPTER = PythonPptxPowerPointHostAdapter()
_DIFF = SlideStateDiffService()


def test_build_path_map_returns_group_and_child(tmp_path):
    from pptx import Presentation
    pptx_path, _ = make_pingpong_pptx(tmp_path)
    prs = Presentation(str(pptx_path))
    path_map = _build_path_map(prs.slides[0].shapes)
    assert "20" in path_map
    assert "20/21" in path_map


def test_child_slide_space_equals_group_plus_local(tmp_path):
    """With chOff=0 and child at local (0,0): child.slide_x == group.slide_x."""
    pptx_path, paths = make_pingpong_pptx(tmp_path)
    state = _ADAPTER.capture_slide_state(pptx_path, slide_number=1)
    shape_map = {s.shape_path: s for s in state.shapes}

    group = shape_map["20"]
    child = shape_map["20/21"]
    assert group.x == _GROUP_X
    assert child.x == _GROUP_X   # local x=0 → slide_x = group_x + 0


def test_parent_group_move_changes_child_slide_space(tmp_path):
    pptx_path, paths = make_pingpong_pptx(tmp_path)
    state_v1 = _ADAPTER.capture_slide_state(pptx_path, slide_number=1)

    child_x_before = next(s.x for s in state_v1.shapes if s.shape_path == "20/21")

    NEW_GROUP_X = _GROUP_X + 500000
    simulate_human_group_move(pptx_path, group_id=_GROUP_ID, new_x=NEW_GROUP_X)

    state_v2 = _ADAPTER.capture_slide_state(pptx_path, slide_number=1)
    child_x_after = next(s.x for s in state_v2.shapes if s.shape_path == "20/21")

    assert child_x_after == child_x_before + 500000


def test_diff_detects_geometry_changed_on_child_after_group_move(tmp_path):
    pptx_path, paths = make_pingpong_pptx(tmp_path)
    state_v1 = _ADAPTER.capture_slide_state(pptx_path, slide_number=1)

    simulate_human_group_move(pptx_path, group_id=_GROUP_ID, new_x=_GROUP_X + 400000)

    state_v2 = _ADAPTER.capture_slide_state(pptx_path, slide_number=1)
    diff = _DIFF.diff(state_v1, state_v2)

    child_path = paths["grouped_child"]
    child_delta = next((d for d in diff.shape_deltas if d.shape_path == child_path), None)
    assert child_delta is not None, "Expected geometry delta on grouped child"
    assert "geometry_changed" in child_delta.delta_categories


def test_formatting_fingerprint_unchanged_after_text_replace(tmp_path):
    """Replacing run text (preserving format) must not change formatting_fingerprint."""
    pptx_path, paths = make_pingpong_pptx(tmp_path)
    state_v1 = _ADAPTER.capture_slide_state(pptx_path, slide_number=1)

    bullet1_path = paths["bullet_1"]
    fp_before = next(s.formatting_fingerprint for s in state_v1.shapes if s.shape_path == bullet1_path)

    state_v1_hash = state_v1.state_hash
    from slidestein.interactive.models import SlidePatchOperation, SlidePatchPlan
    patch = SlidePatchPlan(
        slide_number=1,
        base_state_hash=state_v1_hash,
        operations=[
            SlidePatchOperation(
                operation="replace_text",
                shape_id=12,
                shape_path=bullet1_path,
                field="text",
                value="Replaced via run-preserving patch",
            )
        ],
    )
    result = _ADAPTER.apply_patch(pptx_path, patch)
    assert result.status == "applied"

    fp_after = next(s.formatting_fingerprint for s in result.after_state.shapes if s.shape_path == bullet1_path)
    assert fp_before == fp_after, (
        f"formatting_fingerprint should not change after text replace: {fp_before!r} → {fp_after!r}"
    )
