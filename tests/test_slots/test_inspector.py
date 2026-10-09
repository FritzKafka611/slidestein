"""Tests for SlideShapeInspector — deterministic, uses mock python-pptx objects."""

from __future__ import annotations

from contextlib import contextmanager
from unittest.mock import MagicMock, patch

import pytest

from slidestein.slots.inspector import SlideShapeInspector


# ---------------------------------------------------------------------------
# Mock python-pptx shape factories
# ---------------------------------------------------------------------------

def _mock_text_box(shape_id: int, text: str, left=914400, top=914400,
                   width=4572000, height=685800) -> MagicMock:
    shape = MagicMock()
    shape.shape_id = shape_id
    shape.name = f"TextBox_{shape_id}"
    shape.left = left
    shape.top = top
    shape.width = width
    shape.height = height
    shape.has_text_frame = True
    shape.has_chart = False
    shape.has_table = False
    shape.is_placeholder = False
    # shape_type: AUTO_SHAPE (1)
    from pptx.enum.shapes import MSO_SHAPE_TYPE
    shape.shape_type = MSO_SHAPE_TYPE.AUTO_SHAPE
    tf = MagicMock()
    tf.text = text
    tf.paragraphs = []
    shape.text_frame = tf
    return shape


def _mock_placeholder(shape_id: int, text: str, ph_idx: int, ph_type=15,
                      left=914400, top=457200, width=7315200, height=685800) -> MagicMock:
    shape = MagicMock()
    shape.shape_id = shape_id
    shape.name = f"Title_{shape_id}"
    shape.left = left
    shape.top = top
    shape.width = width
    shape.height = height
    shape.has_text_frame = True
    shape.has_chart = False
    shape.has_table = False
    shape.is_placeholder = True
    from pptx.enum.shapes import MSO_SHAPE_TYPE
    shape.shape_type = MSO_SHAPE_TYPE.AUTO_SHAPE
    tf = MagicMock()
    tf.text = text
    tf.paragraphs = []
    shape.text_frame = tf
    pf = MagicMock()
    pf.type = ph_type
    pf.idx = ph_idx
    shape.placeholder_format = pf
    return shape


def _mock_group(group_id: int, children, left=914400, top=914400,
                width=6858000, height=4114800) -> MagicMock:
    shape = MagicMock()
    shape.shape_id = group_id
    shape.name = f"Group_{group_id}"
    shape.left = left
    shape.top = top
    shape.width = width
    shape.height = height
    shape.has_text_frame = False
    shape.has_chart = False
    shape.has_table = False
    shape.is_placeholder = False
    from pptx.enum.shapes import MSO_SHAPE_TYPE
    shape.shape_type = MSO_SHAPE_TYPE.GROUP
    shape.shapes = children
    return shape


def _mock_table_shape(shape_id: int, rows=3, cols=4) -> MagicMock:
    shape = MagicMock()
    shape.shape_id = shape_id
    shape.name = f"Table_{shape_id}"
    shape.left = 457200
    shape.top = 914400
    shape.width = 6858000
    shape.height = 2743200
    shape.has_text_frame = False
    shape.has_chart = False
    shape.has_table = True
    shape.is_placeholder = False
    from pptx.enum.shapes import MSO_SHAPE_TYPE
    shape.shape_type = MSO_SHAPE_TYPE.AUTO_SHAPE
    tbl = MagicMock()
    tbl.rows = [MagicMock()] * rows
    tbl.columns = [MagicMock()] * cols
    shape.table = tbl
    return shape


def _mock_chart_shape(shape_id: int) -> MagicMock:
    shape = MagicMock()
    shape.shape_id = shape_id
    shape.name = f"Chart_{shape_id}"
    shape.left = 457200
    shape.top = 914400
    shape.width = 6858000
    shape.height = 3429000
    shape.has_text_frame = False
    shape.has_chart = True
    shape.has_table = False
    shape.is_placeholder = False
    from pptx.enum.shapes import MSO_SHAPE_TYPE
    shape.shape_type = MSO_SHAPE_TYPE.AUTO_SHAPE
    return shape


def _mock_slide_record(source_path="C:/deck.pptx", slide_number=1) -> MagicMock:
    record = MagicMock()
    record.source_deck_path = source_path
    record.slide_number = slide_number
    return record


def _patch_pptx(shapes, slide_width=9144000, slide_height=5143500):
    """Context-manager that patches Presentation AND Path.exists to simulate a valid file."""
    mock_slide = MagicMock()
    mock_slide.shapes = shapes
    mock_prs = MagicMock()
    mock_prs.slides = [mock_slide]
    mock_prs.slide_width = slide_width
    mock_prs.slide_height = slide_height

    @contextmanager
    def _ctx():
        with patch("slidestein.slots.inspector.Presentation", return_value=mock_prs), \
             patch("pathlib.Path.exists", return_value=True):
            yield

    return _ctx()


# ---------------------------------------------------------------------------
# Inspector tests
# ---------------------------------------------------------------------------

class TestTextBoxShape:
    def test_text_box_source_kind(self):
        shapes = [_mock_text_box(shape_id=5, text="Hello world")]
        record = _mock_slide_record()
        with _patch_pptx(shapes):
            inspector = SlideShapeInspector()
            descriptors = inspector.inspect(record)
        assert len(descriptors) == 1
        assert descriptors[0].source_kind == "text_box"

    def test_text_box_can_edit_true(self):
        shapes = [_mock_text_box(shape_id=5, text="Hello world")]
        record = _mock_slide_record()
        with _patch_pptx(shapes):
            descriptors = SlideShapeInspector().inspect(record)
        assert descriptors[0].can_edit_text is True

    def test_text_box_text_extracted(self):
        shapes = [_mock_text_box(shape_id=5, text="Alpha Beta")]
        record = _mock_slide_record()
        with _patch_pptx(shapes):
            descriptors = SlideShapeInspector().inspect(record)
        assert descriptors[0].text == "Alpha Beta"
        assert descriptors[0].has_text is True

    def test_text_box_shape_path_is_shape_id_string(self):
        shapes = [_mock_text_box(shape_id=42, text="x")]
        record = _mock_slide_record()
        with _patch_pptx(shapes):
            descriptors = SlideShapeInspector().inspect(record)
        assert descriptors[0].shape_path == "42"

    def test_text_box_parent_shape_path_is_none(self):
        shapes = [_mock_text_box(shape_id=5, text="x")]
        record = _mock_slide_record()
        with _patch_pptx(shapes):
            descriptors = SlideShapeInspector().inspect(record)
        assert descriptors[0].parent_shape_path is None


class TestPlaceholderShape:
    def test_placeholder_source_kind(self):
        shapes = [_mock_placeholder(shape_id=2, text="Slide Title", ph_idx=0)]
        record = _mock_slide_record()
        with _patch_pptx(shapes):
            descriptors = SlideShapeInspector().inspect(record)
        assert descriptors[0].source_kind == "placeholder"

    def test_placeholder_is_placeholder_flag(self):
        shapes = [_mock_placeholder(shape_id=2, text="Title", ph_idx=0)]
        record = _mock_slide_record()
        with _patch_pptx(shapes):
            descriptors = SlideShapeInspector().inspect(record)
        assert descriptors[0].is_placeholder is True

    def test_placeholder_can_edit_true(self):
        shapes = [_mock_placeholder(shape_id=2, text="Title", ph_idx=0)]
        record = _mock_slide_record()
        with _patch_pptx(shapes):
            descriptors = SlideShapeInspector().inspect(record)
        assert descriptors[0].can_edit_text is True

    def test_placeholder_idx_extracted(self):
        shapes = [_mock_placeholder(shape_id=2, text="Title", ph_idx=0)]
        record = _mock_slide_record()
        with _patch_pptx(shapes):
            descriptors = SlideShapeInspector().inspect(record)
        assert descriptors[0].placeholder_idx == 0


class TestGroupShape:
    def test_group_shape_not_editable(self):
        child = _mock_text_box(shape_id=8, text="Child text")
        group = _mock_group(group_id=10, children=[child])
        record = _mock_slide_record()
        with _patch_pptx([group]):
            descriptors = SlideShapeInspector().inspect(record)
        group_desc = next(d for d in descriptors if d.shape_id == 10)
        assert group_desc.can_edit_text is False
        assert group_desc.is_group is True

    def test_group_child_included(self):
        child = _mock_text_box(shape_id=8, text="Child text")
        group = _mock_group(group_id=10, children=[child])
        record = _mock_slide_record()
        with _patch_pptx([group]):
            descriptors = SlideShapeInspector().inspect(record)
        paths = [d.shape_path for d in descriptors]
        assert "10/8" in paths

    def test_group_child_has_parent_path(self):
        child = _mock_text_box(shape_id=8, text="Child text")
        group = _mock_group(group_id=10, children=[child])
        record = _mock_slide_record()
        with _patch_pptx([group]):
            descriptors = SlideShapeInspector().inspect(record)
        child_desc = next(d for d in descriptors if d.shape_path == "10/8")
        assert child_desc.parent_shape_path == "10"

    def test_group_child_can_edit_text(self):
        child = _mock_text_box(shape_id=8, text="Child text")
        group = _mock_group(group_id=10, children=[child])
        record = _mock_slide_record()
        with _patch_pptx([group]):
            descriptors = SlideShapeInspector().inspect(record)
        child_desc = next(d for d in descriptors if d.shape_path == "10/8")
        assert child_desc.can_edit_text is True


class TestTableShape:
    def test_table_source_kind(self):
        shapes = [_mock_table_shape(shape_id=15, rows=3, cols=4)]
        record = _mock_slide_record()
        with _patch_pptx(shapes):
            descriptors = SlideShapeInspector().inspect(record)
        assert descriptors[0].source_kind == "table"

    def test_table_not_editable(self):
        shapes = [_mock_table_shape(shape_id=15)]
        record = _mock_slide_record()
        with _patch_pptx(shapes):
            descriptors = SlideShapeInspector().inspect(record)
        assert descriptors[0].can_edit_text is False

    def test_table_rows_cols_extracted(self):
        shapes = [_mock_table_shape(shape_id=15, rows=3, cols=4)]
        record = _mock_slide_record()
        with _patch_pptx(shapes):
            descriptors = SlideShapeInspector().inspect(record)
        assert descriptors[0].table_rows == 3
        assert descriptors[0].table_columns == 4


class TestChartShape:
    def test_chart_source_kind(self):
        shapes = [_mock_chart_shape(shape_id=20)]
        record = _mock_slide_record()
        with _patch_pptx(shapes):
            descriptors = SlideShapeInspector().inspect(record)
        assert descriptors[0].source_kind == "chart"

    def test_chart_not_editable(self):
        shapes = [_mock_chart_shape(shape_id=20)]
        record = _mock_slide_record()
        with _patch_pptx(shapes):
            descriptors = SlideShapeInspector().inspect(record)
        assert descriptors[0].can_edit_text is False


class TestZOrder:
    def test_z_order_assigned_by_position(self):
        shapes = [
            _mock_text_box(shape_id=1, text="First"),
            _mock_text_box(shape_id=2, text="Second"),
            _mock_text_box(shape_id=3, text="Third"),
        ]
        record = _mock_slide_record()
        with _patch_pptx(shapes):
            descriptors = SlideShapeInspector().inspect(record)
        z_orders = [d.z_order for d in descriptors]
        assert z_orders == [0, 1, 2]


class TestGeometry:
    def test_geometry_in_emu(self):
        # shape at left=914400, top=457200, width=7315200, height=685800
        shapes = [_mock_text_box(shape_id=1, text="x",
                                 left=914400, top=457200, width=7315200, height=685800)]
        record = _mock_slide_record()
        with _patch_pptx(shapes, slide_width=9144000, slide_height=5143500):
            descriptors = SlideShapeInspector().inspect(record)
        d = descriptors[0]
        assert d.x == 914400
        assert d.y == 457200
        assert d.width == 7315200
        assert d.height == 685800

    def test_normalized_ratios_computed(self):
        shapes = [_mock_text_box(shape_id=1, text="x",
                                 left=914400, top=457200, width=7315200, height=685800)]
        record = _mock_slide_record()
        with _patch_pptx(shapes, slide_width=9144000, slide_height=5143500):
            descriptors = SlideShapeInspector().inspect(record)
        d = descriptors[0]
        assert abs(d.x_ratio - 914400 / 9144000) < 1e-4
        assert abs(d.y_ratio - 457200 / 5143500) < 1e-4
        assert abs(d.width_ratio - 7315200 / 9144000) < 1e-4


class TestEmptyTextShape:
    def test_no_text_frame_means_has_text_false(self):
        shape = _mock_chart_shape(shape_id=99)
        record = _mock_slide_record()
        with _patch_pptx([shape]):
            descriptors = SlideShapeInspector().inspect(record)
        assert descriptors[0].has_text is False
        assert descriptors[0].text == ""


class TestFileNotFound:
    def test_missing_source_raises(self):
        # Record points to a path that doesn't exist — inspector must raise
        record = _mock_slide_record(source_path="C:/nonexistent_deck_xyz.pptx")
        inspector = SlideShapeInspector()
        # Do NOT patch Path.exists here — the real file doesn't exist
        with pytest.raises(FileNotFoundError):
            inspector.inspect(record)
