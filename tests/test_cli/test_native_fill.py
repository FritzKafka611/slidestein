"""CLI tests for inspect-text and native-fill commands."""

from __future__ import annotations

from pathlib import Path
from unittest.mock import MagicMock, patch

import pytest
from typer.testing import CliRunner

from slidestein.cli import app
from slidestein.pptx.com_editor import ReplacementResult, ReplacementStatus

runner = CliRunner()


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------


def _make_pptmaster_mock(output_path_ref: list[Path]) -> MagicMock:
    """Return a mock PPTMasterRealAdapter whose extract_slide creates a file."""
    adapter = MagicMock()

    def fake_extract(source_pptx, slide_number, output_path):
        output_path.parent.mkdir(parents=True, exist_ok=True)
        output_path.write_bytes(b"PK_fake_pptx")
        output_path_ref.append(output_path)
        return output_path

    adapter.extract_slide.side_effect = fake_extract
    return adapter


def _verified_result(replacements: dict[str, str]) -> list[ReplacementResult]:
    """Return REPLACED_AND_VERIFIED results for the given replacements."""
    return [
        ReplacementResult(
            requested=k,
            replacement=v,
            status=ReplacementStatus.REPLACED_AND_VERIFIED,
            occurrences_before=1,
            occurrences_after=0,
            matched=True,
            replaced=True,
            verified=True,
        )
        for k, v in replacements.items()
        if k
    ]


def _make_com_editor_mock() -> MagicMock:
    """Return a mock PowerPointComEditor whose replace_text writes a file and returns
    REPLACED_AND_VERIFIED results."""
    editor = MagicMock()

    def fake_replace(source_pptx, replacements, output_path, slide_number=1):
        output_path.parent.mkdir(parents=True, exist_ok=True)
        output_path.write_bytes(b"PK_native_edited")
        return _verified_result(replacements)

    editor.replace_text.side_effect = fake_replace
    return editor


# ---------------------------------------------------------------------------
# inspect-text
# ---------------------------------------------------------------------------


class TestInspectTextCommand:
    def test_missing_source_exits_1(self, tmp_path: Path) -> None:
        result = runner.invoke(app, [
            "inspect-text",
            str(tmp_path / "nonexistent.pptx"),
        ])
        assert result.exit_code == 1
        assert "not found" in result.output.lower()

    def test_non_pptx_exits_1(self, tmp_path: Path) -> None:
        bad = tmp_path / "file.pdf"
        bad.write_bytes(b"%%PDF")
        result = runner.invoke(app, ["inspect-text", str(bad)])
        assert result.exit_code == 1

    def test_out_of_range_slide_exits_1(self, tmp_path: Path) -> None:
        from pptx import Presentation
        prs = Presentation()
        pptx_path = tmp_path / "one_slide.pptx"
        prs.save(str(pptx_path))

        result = runner.invoke(app, [
            "inspect-text", str(pptx_path), "--slide", "99"
        ])
        assert result.exit_code == 1
        assert "out of range" in result.output.lower()

    def test_valid_pptx_exits_0(self, tmp_path: Path) -> None:
        from pptx import Presentation
        from pptx.util import Inches, Pt

        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[5])
        txBox = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1))
        txBox.text_frame.text = "Hello inspect"
        pptx_path = tmp_path / "test.pptx"
        prs.save(str(pptx_path))

        result = runner.invoke(app, ["inspect-text", str(pptx_path)])
        assert result.exit_code == 0
        assert "Hello inspect" in result.output

    def test_shows_shape_id_and_name(self, tmp_path: Path) -> None:
        from pptx import Presentation
        from pptx.util import Inches

        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[5])
        txBox = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1))
        txBox.text_frame.text = "Test content"
        pptx_path = tmp_path / "test.pptx"
        prs.save(str(pptx_path))

        result = runner.invoke(app, ["inspect-text", str(pptx_path)])
        assert result.exit_code == 0
        # Shape ID should appear (numeric)
        assert any(c.isdigit() for c in result.output)

    def test_no_text_shapes_message(self, tmp_path: Path) -> None:
        from pptx import Presentation

        prs = Presentation()
        prs.slides.add_slide(prs.slide_layouts[6])  # blank slide
        pptx_path = tmp_path / "blank.pptx"
        prs.save(str(pptx_path))

        result = runner.invoke(app, ["inspect-text", str(pptx_path)])
        assert result.exit_code == 0
        # Should say no text shapes or show empty table


# ---------------------------------------------------------------------------
# native-fill
# ---------------------------------------------------------------------------


class TestNativeFillCommand:
    def test_missing_source_exits_1(self, tmp_path: Path) -> None:
        result = runner.invoke(app, [
            "native-fill",
            str(tmp_path / "nonexistent.pptx"),
            "3",
            "--replace", "OLD", "NEW",
            "--output", str(tmp_path / "out.pptx"),
        ])
        assert result.exit_code == 1
        assert "not found" in result.output.lower()

    def test_non_pptx_exits_1(self, tmp_path: Path) -> None:
        bad = tmp_path / "file.pdf"
        bad.write_bytes(b"%%PDF")
        result = runner.invoke(app, [
            "native-fill", str(bad), "3",
            "--replace", "OLD", "NEW",
            "--output", str(tmp_path / "out.pptx"),
        ])
        assert result.exit_code == 1

    def test_zero_slide_number_exits_1(self, tmp_path: Path) -> None:
        pptx = tmp_path / "deck.pptx"
        pptx.write_bytes(b"PK")
        result = runner.invoke(app, [
            "native-fill", str(pptx), "0",
            "--replace", "OLD", "NEW",
            "--output", str(tmp_path / "out.pptx"),
        ])
        assert result.exit_code == 1

    def test_success_exits_0(self, tmp_path: Path) -> None:
        pptx = tmp_path / "deck.pptx"
        pptx.write_bytes(b"PK")
        output = tmp_path / "out.pptx"
        output_ref: list[Path] = []

        mock_pptmaster = _make_pptmaster_mock(output_ref)
        mock_editor = _make_com_editor_mock()

        with (
            patch(
                "slidestein.pptx.pptmaster_real_adapter.PPTMasterRealAdapter.from_settings",
                return_value=mock_pptmaster,
            ),
            patch(
                "slidestein.pptx.com_editor.PowerPointComEditor.replace_text",
                side_effect=mock_editor.replace_text,
            ),
            patch(
                "slidestein.pptx.python_pptx_adapter.PythonPptxAdapter.render_preview",
                side_effect=RuntimeError("No renderer"),
            ),
        ):
            result = runner.invoke(app, [
                "native-fill", str(pptx), "3",
                "--replace", "Agenda", "Our Plan",
                "--output", str(output),
            ])

        assert result.exit_code == 0
        mock_pptmaster.extract_slide.assert_called_once()

    def test_replacement_args_passed_correctly(self, tmp_path: Path) -> None:
        pptx = tmp_path / "deck.pptx"
        pptx.write_bytes(b"PK")
        output = tmp_path / "out.pptx"
        captured: list[dict] = []

        def fake_replace(source_pptx, replacements, output_path, slide_number=1):
            captured.append({"replacements": replacements})
            output_path.parent.mkdir(parents=True, exist_ok=True)
            output_path.write_bytes(b"PK")
            return _verified_result(replacements)

        mock_pptmaster = _make_pptmaster_mock([])

        with (
            patch(
                "slidestein.pptx.pptmaster_real_adapter.PPTMasterRealAdapter.from_settings",
                return_value=mock_pptmaster,
            ),
            patch(
                "slidestein.pptx.com_editor.PowerPointComEditor.replace_text",
                side_effect=fake_replace,
            ),
            patch(
                "slidestein.pptx.python_pptx_adapter.PythonPptxAdapter.render_preview",
                side_effect=RuntimeError("No renderer"),
            ),
        ):
            runner.invoke(app, [
                "native-fill", str(pptx), "5",
                "--replace", "Agenda", "Our Plan",
                "--output", str(output),
            ])

        assert len(captured) == 1
        assert captured[0]["replacements"] == {"Agenda": "Our Plan"}

    def test_pptmaster_not_found_exits_1(self, tmp_path: Path) -> None:
        pptx = tmp_path / "deck.pptx"
        pptx.write_bytes(b"PK")

        with patch(
            "slidestein.pptx.pptmaster_real_adapter.PPTMasterRealAdapter.from_settings",
            side_effect=FileNotFoundError("ppt-master venv not found"),
        ):
            result = runner.invoke(app, [
                "native-fill", str(pptx), "3",
                "--replace", "OLD", "NEW",
                "--output", str(tmp_path / "out.pptx"),
            ])

        assert result.exit_code == 1
        assert "ppt-master" in result.output.lower()

    def test_pptmaster_error_exits_1(self, tmp_path: Path) -> None:
        from slidestein.pptx.pptmaster_real_adapter import PPTMasterError

        pptx = tmp_path / "deck.pptx"
        pptx.write_bytes(b"PK")
        mock_pptmaster = MagicMock()
        mock_pptmaster.extract_slide.side_effect = PPTMasterError("svg-to-pptx failed")

        with patch(
            "slidestein.pptx.pptmaster_real_adapter.PPTMasterRealAdapter.from_settings",
            return_value=mock_pptmaster,
        ):
            result = runner.invoke(app, [
                "native-fill", str(pptx), "3",
                "--replace", "OLD", "NEW",
                "--output", str(tmp_path / "out.pptx"),
            ])

        assert result.exit_code == 1
        assert "ppt-master error" in result.output.lower()

    def test_com_error_exits_1(self, tmp_path: Path) -> None:
        from slidestein.pptx.com_editor import PowerPointComEditorError

        pptx = tmp_path / "deck.pptx"
        pptx.write_bytes(b"PK")
        mock_pptmaster = _make_pptmaster_mock([])

        with (
            patch(
                "slidestein.pptx.pptmaster_real_adapter.PPTMasterRealAdapter.from_settings",
                return_value=mock_pptmaster,
            ),
            patch(
                "slidestein.pptx.com_editor.PowerPointComEditor.replace_text",
                side_effect=PowerPointComEditorError("Access denied"),
            ),
        ):
            result = runner.invoke(app, [
                "native-fill", str(pptx), "3",
                "--replace", "OLD", "NEW",
                "--output", str(tmp_path / "out.pptx"),
            ])

        assert result.exit_code == 1
        assert "powerpoint com error" in result.output.lower()

    def test_preview_failure_does_not_fail_command(self, tmp_path: Path) -> None:
        pptx = tmp_path / "deck.pptx"
        pptx.write_bytes(b"PK")
        output = tmp_path / "out.pptx"
        mock_pptmaster = _make_pptmaster_mock([])

        with (
            patch(
                "slidestein.pptx.pptmaster_real_adapter.PPTMasterRealAdapter.from_settings",
                return_value=mock_pptmaster,
            ),
            patch(
                "slidestein.pptx.com_editor.PowerPointComEditor.replace_text",
                side_effect=_make_com_editor_mock().replace_text,
            ),
            patch(
                "slidestein.pptx.python_pptx_adapter.PythonPptxAdapter.render_preview",
                side_effect=RuntimeError("No rendering backend"),
            ),
        ):
            result = runner.invoke(app, [
                "native-fill", str(pptx), "3",
                "--replace", "Agenda", "Our Plan",
                "--output", str(output),
            ])

        assert result.exit_code == 0
        assert "preview skipped" in result.output.lower()

    def test_no_match_warning_shown(self, tmp_path: Path) -> None:
        pptx = tmp_path / "deck.pptx"
        pptx.write_bytes(b"PK")
        output = tmp_path / "out.pptx"
        mock_pptmaster = _make_pptmaster_mock([])

        def fake_replace_no_match(source_pptx, replacements, output_path, slide_number=1):
            output_path.parent.mkdir(parents=True, exist_ok=True)
            output_path.write_bytes(b"PK")
            return [
                ReplacementResult(
                    requested=k,
                    replacement=v,
                    status=ReplacementStatus.NO_MATCH,
                    occurrences_before=0,
                    occurrences_after=0,
                    matched=False,
                    replaced=False,
                    verified=False,
                )
                for k, v in replacements.items()
                if k
            ]

        with (
            patch(
                "slidestein.pptx.pptmaster_real_adapter.PPTMasterRealAdapter.from_settings",
                return_value=mock_pptmaster,
            ),
            patch(
                "slidestein.pptx.com_editor.PowerPointComEditor.replace_text",
                side_effect=fake_replace_no_match,
            ),
            patch(
                "slidestein.pptx.python_pptx_adapter.PythonPptxAdapter.render_preview",
                side_effect=RuntimeError("No renderer"),
            ),
        ):
            result = runner.invoke(app, [
                "native-fill", str(pptx), "3",
                "--replace", "NonExistent", "Something",
                "--output", str(output),
            ])

        assert result.exit_code == 0
        assert "warning" in result.output.lower()
        assert "not found" in result.output.lower()

    def test_error_status_exits_1(self, tmp_path: Path) -> None:
        """CLI must exit non-zero when replace_text returns ERROR status."""
        pptx = tmp_path / "deck.pptx"
        pptx.write_bytes(b"PK")
        output = tmp_path / "out.pptx"
        mock_pptmaster = _make_pptmaster_mock([])

        def fake_replace_error(source_pptx, replacements, output_path, slide_number=1):
            output_path.parent.mkdir(parents=True, exist_ok=True)
            output_path.write_bytes(b"PK")
            return [
                ReplacementResult(
                    requested=k,
                    replacement=v,
                    status=ReplacementStatus.ERROR,
                    occurrences_before=1,
                    occurrences_after=1,
                    matched=True,
                    replaced=True,
                    verified=False,
                    error="SaveAs succeeded but text unchanged",
                )
                for k, v in replacements.items()
                if k
            ]

        with (
            patch(
                "slidestein.pptx.pptmaster_real_adapter.PPTMasterRealAdapter.from_settings",
                return_value=mock_pptmaster,
            ),
            patch(
                "slidestein.pptx.com_editor.PowerPointComEditor.replace_text",
                side_effect=fake_replace_error,
            ),
            patch(
                "slidestein.pptx.python_pptx_adapter.PythonPptxAdapter.render_preview",
                side_effect=RuntimeError("No renderer"),
            ),
        ):
            result = runner.invoke(app, [
                "native-fill", str(pptx), "3",
                "--replace", "Agenda", "Our Plan",
                "--output", str(output),
            ])

        assert result.exit_code == 1
        assert "error" in result.output.lower()
        assert "verification" in result.output.lower()
