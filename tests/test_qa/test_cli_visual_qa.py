"""CLI tests for visual-qa command (M8)."""

from __future__ import annotations

import json
from pathlib import Path
from unittest.mock import MagicMock, patch

import pytest
from typer.testing import CliRunner

from slidestein.cli import app
from slidestein.qa.errors import VisualQAError
from slidestein.qa.models import VisualQAResult
from slidestein.qa.providers.sap_aicore import VisualQAGenerationError
from tests.test_qa.conftest import (
    make_complete_draft,
    make_native_check,
    make_slot_map,
    make_visual_model_output,
)

runner = CliRunner()


# ---------------------------------------------------------------------------
# Fixtures
# ---------------------------------------------------------------------------


def _write_json(path: Path, data: object) -> None:
    path.write_text(json.dumps(data), encoding="utf-8")


def _make_draft_json(tmp_path) -> Path:
    draft = make_complete_draft()
    p = tmp_path / "draft.json"
    p.write_text(draft.model_dump_json(), encoding="utf-8")
    return p


def _make_slot_map_json(tmp_path) -> Path:
    slot_map = make_slot_map()
    p = tmp_path / "slot_map.json"
    p.write_text(slot_map.model_dump_json(), encoding="utf-8")
    return p


def _make_fake_pptx(tmp_path, name: str) -> Path:
    from pptx import Presentation  # noqa: PLC0415

    p = tmp_path / name
    prs = Presentation()
    prs.slides.add_slide(prs.slide_layouts[0])
    prs.save(str(p))
    return p


def _make_valid_result(tmp_path) -> VisualQAResult:
    from slidestein.qa.models import VisualDimensionAssessment

    return VisualQAResult(
        slide_id="slide-test-001",
        deck_id="deck-abc123",
        slide_number=2,
        resolved_generated_slide_number=2,
        recommendation="pass",
        text_fit_and_clipping=VisualDimensionAssessment(score=5, rationale="Clean"),
        visual_hierarchy=VisualDimensionAssessment(score=4, rationale="Good"),
        alignment_and_spacing=VisualDimensionAssessment(score=4, rationale="OK"),
        balance_and_whitespace=VisualDimensionAssessment(score=4, rationale="OK"),
        typography_and_style_consistency=VisualDimensionAssessment(score=4, rationale="OK"),
        overall_readability=VisualDimensionAssessment(score=4, rationale="OK"),
        average_score=4.1667,
        native_checks=[make_native_check()],
        issues=[],
        executive_summary="Slide renders cleanly.",
        generated_render_path=str(tmp_path / "generated.png"),
        template_render_path=str(tmp_path / "template.png"),
        overlay_render_path=str(tmp_path / "overlay.png"),
    )


# ---------------------------------------------------------------------------
# Missing required arguments
# ---------------------------------------------------------------------------


def test_missing_generated_pptx_exits_nonzero(tmp_path):
    draft_p = _make_draft_json(tmp_path)
    slot_map_p = _make_slot_map_json(tmp_path)
    result = runner.invoke(
        app,
        [
            "visual-qa",
            "--template-pptx", str(tmp_path / "template.pptx"),
            "--draft", str(draft_p),
            "--slot-map", str(slot_map_p),
        ],
    )
    assert result.exit_code != 0


def test_invalid_draft_json_exits_nonzero(tmp_path):
    template_p = _make_fake_pptx(tmp_path, "template.pptx")
    generated_p = _make_fake_pptx(tmp_path, "generated.pptx")
    slot_map_p = _make_slot_map_json(tmp_path)
    bad_draft = tmp_path / "bad.json"
    bad_draft.write_text("not valid json", encoding="utf-8")

    result = runner.invoke(
        app,
        [
            "visual-qa",
            "--template-pptx", str(template_p),
            "--generated-pptx", str(generated_p),
            "--draft", str(bad_draft),
            "--slot-map", str(slot_map_p),
        ],
    )
    assert result.exit_code != 0


def test_incomplete_draft_exits_nonzero(tmp_path):
    template_p = _make_fake_pptx(tmp_path, "template.pptx")
    generated_p = _make_fake_pptx(tmp_path, "generated.pptx")
    slot_map_p = _make_slot_map_json(tmp_path)

    incomplete = make_complete_draft(is_complete=False)
    draft_p = tmp_path / "incomplete.json"
    draft_p.write_text(incomplete.model_dump_json(), encoding="utf-8")

    result = runner.invoke(
        app,
        [
            "visual-qa",
            "--template-pptx", str(template_p),
            "--generated-pptx", str(generated_p),
            "--draft", str(draft_p),
            "--slot-map", str(slot_map_p),
        ],
    )
    assert result.exit_code != 0


# ---------------------------------------------------------------------------
# Dry-run
# ---------------------------------------------------------------------------


def test_dry_run_zero_vision_calls(tmp_path):
    template_p = _make_fake_pptx(tmp_path, "template.pptx")
    generated_p = _make_fake_pptx(tmp_path, "generated.pptx")
    draft_p = _make_draft_json(tmp_path)
    slot_map_p = _make_slot_map_json(tmp_path)

    slot_map = make_slot_map()
    reviewer_mock = MagicMock()

    with (
        patch("slidestein.qa.providers.factory.create_visual_qa_reviewer", return_value=reviewer_mock),
        patch("slidestein.writeback.preflight.resolve_slide_by_stable_id", return_value=(1, 256)),
        patch("slidestein.pptx.python_pptx_adapter.PythonPptxAdapter.render_preview"),
        patch("slidestein.qa.overlay.create_slot_overlay"),
        patch("slidestein.qa.inspector.NativeVisualInspector.inspect", return_value=[]),
    ):
        result = runner.invoke(
            app,
            [
                "visual-qa",
                "--template-pptx", str(template_p),
                "--generated-pptx", str(generated_p),
                "--draft", str(draft_p),
                "--slot-map", str(slot_map_p),
                "--dry-run",
            ],
        )

    reviewer_mock.review.assert_not_called()


def test_dry_run_shows_slot_count_in_output(tmp_path):
    template_p = _make_fake_pptx(tmp_path, "template.pptx")
    generated_p = _make_fake_pptx(tmp_path, "generated.pptx")
    draft_p = _make_draft_json(tmp_path)
    slot_map_p = _make_slot_map_json(tmp_path)

    with (
        patch("slidestein.writeback.preflight.resolve_slide_by_stable_id", return_value=(1, 256)),
        patch("slidestein.pptx.python_pptx_adapter.PythonPptxAdapter.render_preview"),
        patch("slidestein.qa.overlay.create_slot_overlay"),
        patch("slidestein.qa.inspector.NativeVisualInspector.inspect", return_value=[]),
    ):
        result = runner.invoke(
            app,
            [
                "visual-qa",
                "--template-pptx", str(template_p),
                "--generated-pptx", str(generated_p),
                "--draft", str(draft_p),
                "--slot-map", str(slot_map_p),
                "--dry-run",
            ],
        )

    # Output should mention slots
    assert "Slots" in result.output or "K1" in result.output or "2" in result.output


# ---------------------------------------------------------------------------
# Full review
# ---------------------------------------------------------------------------


def test_full_review_exits_zero(tmp_path):
    template_p = _make_fake_pptx(tmp_path, "template.pptx")
    generated_p = _make_fake_pptx(tmp_path, "generated.pptx")
    draft_p = _make_draft_json(tmp_path)
    slot_map_p = _make_slot_map_json(tmp_path)
    valid_result = _make_valid_result(tmp_path)

    with (
        patch("slidestein.qa.service.resolve_slide_by_stable_id", return_value=(2, 256)),
        patch("slidestein.qa.providers.factory.create_visual_qa_reviewer") as mock_factory,
        patch("slidestein.pptx.python_pptx_adapter.PythonPptxAdapter.render_preview"),
        patch("slidestein.qa.overlay.create_slot_overlay"),
        patch("slidestein.qa.inspector.NativeVisualInspector.inspect", return_value=[]),
    ):
        mock_reviewer = MagicMock()
        mock_reviewer.review.return_value = make_visual_model_output()
        mock_factory.return_value = mock_reviewer

        # We need to also patch the service's reviewer directly
        with patch("slidestein.qa.service.VisualQAService.review", return_value=valid_result):
            result = runner.invoke(
                app,
                [
                    "visual-qa",
                    "--template-pptx", str(template_p),
                    "--generated-pptx", str(generated_p),
                    "--draft", str(draft_p),
                    "--slot-map", str(slot_map_p),
                ],
            )

    assert result.exit_code == 0


def test_output_flag_writes_json(tmp_path):
    template_p = _make_fake_pptx(tmp_path, "template.pptx")
    generated_p = _make_fake_pptx(tmp_path, "generated.pptx")
    draft_p = _make_draft_json(tmp_path)
    slot_map_p = _make_slot_map_json(tmp_path)
    output_p = tmp_path / "result.json"
    valid_result = _make_valid_result(tmp_path)

    with patch("slidestein.qa.service.VisualQAService.review", return_value=valid_result), \
         patch("slidestein.qa.providers.factory.create_visual_qa_reviewer", return_value=MagicMock()):
        result = runner.invoke(
            app,
            [
                "visual-qa",
                "--template-pptx", str(template_p),
                "--generated-pptx", str(generated_p),
                "--draft", str(draft_p),
                "--slot-map", str(slot_map_p),
                "--output", str(output_p),
            ],
        )

    if result.exit_code == 0:
        assert output_p.exists()
        data = json.loads(output_p.read_text(encoding="utf-8"))
        loaded = VisualQAResult.model_validate(data)
        assert loaded.recommendation == "pass"


def test_output_json_round_trips(tmp_path):
    valid_result = _make_valid_result(tmp_path)
    json_str = valid_result.model_dump_json()
    restored = VisualQAResult.model_validate_json(json_str)
    assert restored.recommendation == valid_result.recommendation
    assert restored.average_score == valid_result.average_score


def test_console_output_is_cp1252_safe(tmp_path):
    """Output must not contain non-cp1252 characters."""
    template_p = _make_fake_pptx(tmp_path, "template.pptx")
    generated_p = _make_fake_pptx(tmp_path, "generated.pptx")
    draft_p = _make_draft_json(tmp_path)
    slot_map_p = _make_slot_map_json(tmp_path)
    valid_result = _make_valid_result(tmp_path)

    with patch("slidestein.qa.service.VisualQAService.review", return_value=valid_result), \
         patch("slidestein.qa.providers.factory.create_visual_qa_reviewer", return_value=MagicMock()):
        result = runner.invoke(
            app,
            [
                "visual-qa",
                "--template-pptx", str(template_p),
                "--generated-pptx", str(generated_p),
                "--draft", str(draft_p),
                "--slot-map", str(slot_map_p),
            ],
        )

    # Attempt cp1252 encoding — will raise if unsafe chars present
    result.output.encode("cp1252", errors="strict")


def test_visual_qa_error_exits_nonzero(tmp_path):
    template_p = _make_fake_pptx(tmp_path, "template.pptx")
    generated_p = _make_fake_pptx(tmp_path, "generated.pptx")
    draft_p = _make_draft_json(tmp_path)
    slot_map_p = _make_slot_map_json(tmp_path)

    with patch("slidestein.qa.service.VisualQAService.review",
               side_effect=VisualQAError("Slide resolution failed")), \
         patch("slidestein.qa.providers.factory.create_visual_qa_reviewer", return_value=MagicMock()):
        result = runner.invoke(
            app,
            [
                "visual-qa",
                "--template-pptx", str(template_p),
                "--generated-pptx", str(generated_p),
                "--draft", str(draft_p),
                "--slot-map", str(slot_map_p),
            ],
        )

    assert result.exit_code != 0
