"""Tests for slidestein.identity.slide_identity."""

from __future__ import annotations

import io
import shutil
import struct
import uuid
import zipfile
import zlib
from pathlib import Path

import pytest
from pptx import Presentation
from pptx.util import Emu, Pt

from slidestein.identity.slide_identity import (
    compute_content_fingerprint,
    compute_deck_fingerprint,
    compute_input_fingerprint,
    compute_structure_fingerprint,
    extract_slide_identities,
    make_stable_slide_id,
)


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------


def _make_pptx(path: Path, n_slides: int = 3, add_text: bool = True) -> Path:
    prs = Presentation()
    layout = prs.slide_layouts[0]
    for i in range(1, n_slides + 1):
        slide = prs.slides.add_slide(layout)
        if add_text:
            for ph in slide.placeholders:
                if ph.placeholder_format.idx == 0:
                    ph.text = f"Slide {i} Title"
                elif ph.placeholder_format.idx == 1:
                    ph.text = f"Subtitle {i}"
    prs.save(str(path))
    return path


def _make_minimal_png(r: int = 128, g: int = 128, b: int = 128) -> bytes:
    """Return a 1×1 pixel RGB PNG with the given colour."""
    def chunk(tag: bytes, data: bytes) -> bytes:
        crc = zlib.crc32(tag + data) & 0xFFFFFFFF
        return struct.pack(">I", len(data)) + tag + data + struct.pack(">I", crc)

    ihdr_data = struct.pack(">IIBBBBB", 1, 1, 8, 2, 0, 0, 0)
    raw_row   = bytes([0, r, g, b])  # filter byte (none) + RGB
    return (
        b"\x89PNG\r\n\x1a\n"
        + chunk(b"IHDR", ihdr_data)
        + chunk(b"IDAT", zlib.compress(raw_row))
        + chunk(b"IEND", b"")
    )


def _make_pptx_with_image(path: Path, png_bytes: bytes) -> None:
    """Save a 1-slide PPTX with a single embedded picture."""
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank layout
    img_stream = io.BytesIO(png_bytes)
    slide.shapes.add_picture(img_stream, Emu(914400), Emu(914400), Emu(914400), Emu(914400))
    prs.save(str(path))


def _replace_image_in_pptx(src_path: Path, dest_path: Path, new_png: bytes) -> None:
    """Copy src_path → dest_path, replacing the first image in ppt/media/ with new_png."""
    with zipfile.ZipFile(src_path, "r") as z:
        entries = {name: z.read(name) for name in z.namelist()}

    media_names = sorted(
        n for n in entries if n.startswith("ppt/media/")
    )
    if not media_names:
        raise ValueError("No media found in PPTX")

    entries[media_names[0]] = new_png

    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        for name, data in entries.items():
            z.writestr(name, data)
    dest_path.write_bytes(buf.getvalue())


# ---------------------------------------------------------------------------
# extract_slide_identities
# ---------------------------------------------------------------------------


class TestExtractSlideIdentities:
    def test_count_matches_slide_count(self, three_slide_pptx):
        ids = extract_slide_identities(three_slide_pptx)
        assert len(ids) == 3

    def test_native_ids_are_positive_integers(self, three_slide_pptx):
        ids = extract_slide_identities(three_slide_pptx)
        for si in ids:
            assert isinstance(si.native_slide_id, int)
            assert si.native_slide_id > 0

    def test_native_ids_are_unique(self, three_slide_pptx):
        ids = extract_slide_identities(three_slide_pptx)
        native_ids = [si.native_slide_id for si in ids]
        assert len(native_ids) == len(set(native_ids))

    def test_slide_numbers_are_one_indexed(self, three_slide_pptx):
        ids = extract_slide_identities(three_slide_pptx)
        numbers = [si.slide_number for si in ids]
        assert numbers == [1, 2, 3]

    def test_slide_part_is_relative_path(self, three_slide_pptx):
        ids = extract_slide_identities(three_slide_pptx)
        for si in ids:
            assert si.slide_part.startswith("slides/")
            assert si.slide_part.endswith(".xml")

    def test_deterministic(self, three_slide_pptx):
        ids1 = extract_slide_identities(three_slide_pptx)
        ids2 = extract_slide_identities(three_slide_pptx)
        assert [(si.slide_number, si.native_slide_id) for si in ids1] == [
            (si.slide_number, si.native_slide_id) for si in ids2
        ]


class TestSlideIdStability:
    def test_stable_after_text_change(self, three_slide_pptx, tmp_path):
        """Native slide IDs survive a text-only edit."""
        ids_before = extract_slide_identities(three_slide_pptx)

        edited = tmp_path / "edited.pptx"
        shutil.copy(three_slide_pptx, edited)
        prs = Presentation(str(edited))
        for ph in prs.slides[0].placeholders:
            if ph.placeholder_format.idx == 0:
                ph.text = "Completely Different Title"
        prs.save(str(edited))

        ids_after = extract_slide_identities(edited)
        # Same native IDs for each slide position
        assert [si.native_slide_id for si in ids_before] == [
            si.native_slide_id for si in ids_after
        ]

    def test_different_slides_have_different_native_ids(self, three_slide_pptx):
        ids = extract_slide_identities(three_slide_pptx)
        native_ids = {si.native_slide_id for si in ids}
        assert len(native_ids) == 3


# ---------------------------------------------------------------------------
# make_stable_slide_id
# ---------------------------------------------------------------------------


class TestMakeStableSlideId:
    def test_is_valid_uuid(self):
        sid = make_stable_slide_id("deck-abc", 256)
        parsed = uuid.UUID(sid)
        assert parsed.version == 5

    def test_deterministic_for_same_inputs(self):
        a = make_stable_slide_id("deck-xyz", 300)
        b = make_stable_slide_id("deck-xyz", 300)
        assert a == b

    def test_different_for_different_native_ids(self):
        a = make_stable_slide_id("same-deck", 100)
        b = make_stable_slide_id("same-deck", 200)
        assert a != b

    def test_different_for_different_deck_ids(self):
        a = make_stable_slide_id("deck-A", 100)
        b = make_stable_slide_id("deck-B", 100)
        assert a != b


# ---------------------------------------------------------------------------
# compute_content_fingerprint
# ---------------------------------------------------------------------------


class TestContentFingerprint:
    def test_stable_for_same_file(self, three_slide_pptx):
        ids = extract_slide_identities(three_slide_pptx)
        fp1 = compute_content_fingerprint(three_slide_pptx, ids[0].slide_part)
        fp2 = compute_content_fingerprint(three_slide_pptx, ids[0].slide_part)
        assert fp1 == fp2

    def test_different_slides_have_different_fingerprints(self, three_slide_pptx):
        ids = extract_slide_identities(three_slide_pptx)
        fps = [compute_content_fingerprint(three_slide_pptx, si.slide_part) for si in ids]
        assert len(fps) == len(set(fps))

    def test_changes_after_text_edit(self, three_slide_pptx, tmp_path):
        ids = extract_slide_identities(three_slide_pptx)
        fp_before = compute_content_fingerprint(three_slide_pptx, ids[0].slide_part)

        edited = tmp_path / "edited.pptx"
        shutil.copy(three_slide_pptx, edited)
        prs = Presentation(str(edited))
        for ph in prs.slides[0].placeholders:
            if ph.placeholder_format.idx == 0:
                ph.text = "Changed Title XYZ"
        prs.save(str(edited))

        ids_after = extract_slide_identities(edited)
        fp_after = compute_content_fingerprint(edited, ids_after[0].slide_part)
        assert fp_before != fp_after

    def test_returns_64_char_hex(self, three_slide_pptx):
        ids = extract_slide_identities(three_slide_pptx)
        fp = compute_content_fingerprint(three_slide_pptx, ids[0].slide_part)
        assert len(fp) == 64
        assert all(c in "0123456789abcdef" for c in fp)


# ---------------------------------------------------------------------------
# compute_structure_fingerprint
# ---------------------------------------------------------------------------


class TestStructureFingerprint:
    def test_stable_for_same_file(self, three_slide_pptx):
        fp1 = compute_structure_fingerprint(three_slide_pptx, 1)
        fp2 = compute_structure_fingerprint(three_slide_pptx, 1)
        assert fp1 == fp2

    def test_unchanged_after_text_edit(self, three_slide_pptx, tmp_path):
        fp_before = compute_structure_fingerprint(three_slide_pptx, 1)

        edited = tmp_path / "edited.pptx"
        shutil.copy(three_slide_pptx, edited)
        prs = Presentation(str(edited))
        for ph in prs.slides[0].placeholders:
            if ph.placeholder_format.idx == 0:
                ph.text = "New Text That Changes Content"
        prs.save(str(edited))

        fp_after = compute_structure_fingerprint(edited, 1)
        assert fp_before == fp_after

    def test_returns_64_char_hex(self, three_slide_pptx):
        fp = compute_structure_fingerprint(three_slide_pptx, 1)
        assert len(fp) == 64
        assert all(c in "0123456789abcdef" for c in fp)


# ---------------------------------------------------------------------------
# compute_deck_fingerprint
# ---------------------------------------------------------------------------


class TestDeckFingerprint:
    def test_returns_64_char_hex(self, three_slide_pptx):
        fp = compute_deck_fingerprint(three_slide_pptx)
        assert len(fp) == 64
        assert all(c in "0123456789abcdef" for c in fp)

    def test_different_files_have_different_fingerprints(self, tmp_path):
        p1 = _make_pptx(tmp_path / "a.pptx", n_slides=1)
        p2 = _make_pptx(tmp_path / "b.pptx", n_slides=2)
        assert compute_deck_fingerprint(p1) != compute_deck_fingerprint(p2)

    def test_deterministic(self, three_slide_pptx):
        assert compute_deck_fingerprint(three_slide_pptx) == compute_deck_fingerprint(
            three_slide_pptx
        )


# ---------------------------------------------------------------------------
# compute_input_fingerprint
# ---------------------------------------------------------------------------


class TestInputFingerprint:
    def test_deterministic(self):
        fp1 = compute_input_fingerprint("slide-abc", "cfp", "sfp", "1.0", "1.0")
        fp2 = compute_input_fingerprint("slide-abc", "cfp", "sfp", "1.0", "1.0")
        assert fp1 == fp2

    def test_changes_with_content_fingerprint(self):
        a = compute_input_fingerprint("slide-abc", "cfp1", "sfp", "1.0", "1.0")
        b = compute_input_fingerprint("slide-abc", "cfp2", "sfp", "1.0", "1.0")
        assert a != b

    def test_changes_with_classification_version(self):
        a = compute_input_fingerprint("slide-abc", "cfp", "sfp", "1.0", "1.0")
        b = compute_input_fingerprint("slide-abc", "cfp", "sfp", "2.0", "1.0")
        assert a != b

    def test_returns_64_char_hex(self):
        fp = compute_input_fingerprint("slide", "c", "s", "1.0", "1.0")
        assert len(fp) == 64
        assert all(c in "0123456789abcdef" for c in fp)


# ---------------------------------------------------------------------------
# M3.4.1 Fix 3 — strengthened content_fingerprint (image replacement)
# ---------------------------------------------------------------------------


class TestContentFingerprintStrengthened:
    """Verify that compute_content_fingerprint detects image swaps and that
    text-only slides remain unchanged (backward compat for existing records)."""

    def test_fingerprint_changes_after_image_replacement(self, tmp_path):
        """Replacing image bytes with different content must change the fingerprint."""
        png_red   = _make_minimal_png(r=255, g=0, b=0)
        png_blue  = _make_minimal_png(r=0, g=0, b=255)

        src_path  = tmp_path / "with_red.pptx"
        dest_path = tmp_path / "with_blue.pptx"

        _make_pptx_with_image(src_path, png_red)
        _replace_image_in_pptx(src_path, dest_path, png_blue)

        ids_src  = extract_slide_identities(src_path)
        ids_dest = extract_slide_identities(dest_path)

        fp_src  = compute_content_fingerprint(src_path,  ids_src[0].slide_part)
        fp_dest = compute_content_fingerprint(dest_path, ids_dest[0].slide_part)

        assert fp_src != fp_dest, "fingerprint must change when image bytes change"

    def test_fingerprint_stable_for_same_image(self, tmp_path):
        """Two PPTX files with the same image bytes produce the same fingerprint."""
        png = _make_minimal_png(r=100, g=150, b=200)
        path1 = tmp_path / "a.pptx"
        path2 = tmp_path / "b.pptx"
        _make_pptx_with_image(path1, png)
        _make_pptx_with_image(path2, png)

        ids1 = extract_slide_identities(path1)
        ids2 = extract_slide_identities(path2)

        fp1 = compute_content_fingerprint(path1, ids1[0].slide_part)
        fp2 = compute_content_fingerprint(path2, ids2[0].slide_part)

        assert fp1 == fp2

    def test_text_only_slide_fingerprint_unchanged_vs_v1(self, three_slide_pptx):
        """For text-only slides (no images/charts), fingerprint equals SHA-256 of
        slide XML — same as the original v1 algorithm.  Existing records stay valid."""
        import hashlib
        import zipfile as zf

        ids = extract_slide_identities(three_slide_pptx)
        slide_part = ids[0].slide_part

        # v1 result: raw SHA-256 of slide XML only.
        with zf.ZipFile(three_slide_pptx, "r") as z:
            xml_bytes = z.read(f"ppt/{slide_part}")
        v1_fp = hashlib.sha256(xml_bytes).hexdigest()

        # v2 result: should equal v1 when no renderable relationships exist.
        v2_fp = compute_content_fingerprint(three_slide_pptx, slide_part)

        assert v1_fp == v2_fp, "text-only slide fingerprint must not change in v2"

    def test_unrelated_slide_edit_does_not_change_other_fingerprint(
        self, three_slide_pptx, tmp_path
    ):
        """An edit to slide 2 must not change slide 1's content_fingerprint."""
        ids_before = extract_slide_identities(three_slide_pptx)
        fp_slide1_before = compute_content_fingerprint(
            three_slide_pptx, ids_before[0].slide_part
        )

        edited = tmp_path / "edited.pptx"
        shutil.copy(three_slide_pptx, edited)
        prs = Presentation(str(edited))
        # Edit only slide 2.
        for ph in prs.slides[1].placeholders:
            if ph.placeholder_format.idx == 0:
                ph.text = "Completely Different Slide 2 Title"
        prs.save(str(edited))

        ids_after = extract_slide_identities(edited)
        fp_slide1_after = compute_content_fingerprint(
            edited, ids_after[0].slide_part
        )

        assert fp_slide1_before == fp_slide1_after, (
            "editing slide 2 must not change slide 1's content fingerprint"
        )
