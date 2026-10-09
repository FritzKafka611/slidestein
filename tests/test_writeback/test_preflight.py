"""Tests for writeback/preflight.py — 9-step preflight validation.

All python-pptx, fingerprint, and identity calls are mocked.
No real PPTX files, no COM, no LLM.
"""

from __future__ import annotations

from contextlib import contextmanager
from pathlib import Path
from typing import Iterator
from unittest.mock import MagicMock, patch

import pytest
from pydantic import ValidationError

from slidestein.drafting.models import (
    SlideContentDraft,
    SlotDraftAction,
    SlotDraftAssignment,
)
from slidestein.slots.models import SlotCapacity, TemplateSlot, TemplateSlotMap
from slidestein.slots.roles import SlotRole
from slidestein.writeback.errors import PowerPointWritebackError
from slidestein.writeback.models import PowerPointWritebackRequest, WritebackPlan
from slidestein.writeback.preflight import preflight
from slidestein.writeback.versions import WRITEBACK_SCHEMA_VERSION


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------

_SLIDE_ID = "slide-ab12"
_DECK_ID = "deck-xyz"
_SLIDE_NUMBER = 2
_NATIVE_SLIDE_ID = 256
_SHAPE_ID = 7
_SHAPE_PATH = "7"


def _capacity() -> SlotCapacity:
    return SlotCapacity(
        max_characters_estimate=100,
        max_lines_estimate=2,
        current_character_count=10,
        current_line_count=1,
        relative_capacity="medium",
    )


def _slot(
    slot_id: str = "s1",
    shape_id: int = _SHAPE_ID,
    shape_path: str = _SHAPE_PATH,
) -> TemplateSlot:
    return TemplateSlot(
        slot_id=slot_id,
        shape_id=shape_id,
        shape_path=shape_path,
        slot_role=SlotRole.TITLE,
        semantic_label="Title",
        editable=True,
        confidence=0.9,
        current_text="Old title",
        geometry={"x_ratio": 0.05, "y_ratio": 0.04, "width_ratio": 0.9, "height_ratio": 0.1},
        capacity=_capacity(),
    )


def _slot_map(
    slots: list[TemplateSlot] | None = None,
    slide_id: str = _SLIDE_ID,
    deck_id: str | None = _DECK_ID,
    slide_number: int = _SLIDE_NUMBER,
) -> TemplateSlotMap:
    return TemplateSlotMap(
        slide_id=slide_id,
        deck_id=deck_id,
        slide_number=slide_number,
        slide_width=9144000,
        slide_height=5143500,
        slots=slots or [_slot()],
        non_editable_elements=[],
        unsupported_elements=[],
        groups=[],
        analysis_summary="test",
        slot_analysis_input_fingerprint="fp-test",
    )


def _assignment(
    slot_id: str = "s1",
    shape_id: int = _SHAPE_ID,
    shape_path: str = _SHAPE_PATH,
    action: SlotDraftAction = SlotDraftAction.REPLACE,
    text: str | None = "New title",
) -> SlotDraftAssignment:
    return SlotDraftAssignment(
        slot_id=slot_id,
        shape_id=shape_id,
        shape_path=shape_path,
        slot_role=SlotRole.TITLE,
        semantic_label="Title",
        action=action,
        text=text,
        capacity_characters=100,
        capacity_utilization=0.1,
        max_lines_estimate=2,
    )


def _draft(
    assignments: list[SlotDraftAssignment] | None = None,
    slide_id: str = _SLIDE_ID,
    deck_id: str | None = _DECK_ID,
    slide_number: int = _SLIDE_NUMBER,
    is_complete: bool = True,
) -> SlideContentDraft:
    return SlideContentDraft(
        slide_id=slide_id,
        deck_id=deck_id,
        slide_number=slide_number,
        brief_key_message="Test message.",
        assignments=assignments or [_assignment()],
        drafting_summary="Test.",
        is_complete=is_complete,
    )


def _request(
    draft: SlideContentDraft | None = None,
    slot_map: TemplateSlotMap | None = None,
    source_pptx: Path = Path("/src/deck.pptx"),
    output_pptx: Path = Path("/out/deck.pptx"),
    overwrite: bool = True,
) -> PowerPointWritebackRequest:
    return PowerPointWritebackRequest(
        draft=draft or _draft(),
        slot_map=slot_map or _slot_map(),
        source_pptx=source_pptx,
        output_pptx=output_pptx,
        overwrite=overwrite,
    )


def _mock_pptx_slide(
    shape_id: int = _SHAPE_ID,
    shape_path_key: str = _SHAPE_PATH,
    has_text_frame: bool = True,
) -> MagicMock:
    """Return a mock slide with one shape at the given path."""
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    shape = MagicMock()
    shape.shape_id = shape_id
    shape.shape_type = MSO_SHAPE_TYPE.AUTO_SHAPE
    shape.has_text_frame = has_text_frame

    slide = MagicMock()
    slide.shapes = [shape]
    return slide


@contextmanager
def _hardening_patches(
    fp: str = "aabbccdd",
    resolved_slide_number: int = _SLIDE_NUMBER,
    native_slide_id: int = _NATIVE_SLIDE_ID,
) -> Iterator[None]:
    """Context manager that patches all step-5b/6/8 functions.

    step-5b: resolve_slide_by_stable_id
    step-6:  check_source_currentness (no-op)
    step-8:  compute_structure_fingerprint, compute_format_fingerprint,
             compute_deck_structure_snapshot
    """
    with (
        patch(
            "slidestein.writeback.preflight.resolve_slide_by_stable_id",
            return_value=(resolved_slide_number, native_slide_id),
        ),
        patch(
            "slidestein.writeback.preflight.check_source_currentness",
        ),
        patch(
            "slidestein.identity.slide_identity.compute_structure_fingerprint",
            return_value=fp,
        ),
        patch(
            "slidestein.writeback.verification.compute_format_fingerprint",
            return_value="fmt-fp",
        ),
        patch(
            "slidestein.writeback.verification.compute_deck_structure_snapshot",
            return_value=([native_slide_id], {str(native_slide_id): fp}),
        ),
    ):
        yield


def _patched_preflight(
    req: PowerPointWritebackRequest,
    n_slides: int = 5,
    slide: MagicMock | None = None,
    fp: str = "aabbccdd",
) -> WritebackPlan:
    """Run preflight with mocked python-pptx and all hardening patches."""
    if slide is None:
        slide = _mock_pptx_slide()

    mock_prs = MagicMock()
    mock_prs.slides = [MagicMock()] * n_slides
    mock_prs.slides[req.slot_map.slide_number - 1] = slide

    with (
        patch("slidestein.writeback.preflight.Presentation", return_value=mock_prs),
        _hardening_patches(fp=fp, resolved_slide_number=req.slot_map.slide_number),
    ):
        return preflight(req)


# ---------------------------------------------------------------------------
# Step 1: Completeness gate
# ---------------------------------------------------------------------------


def test_incomplete_draft_rejected() -> None:
    incomplete_assign = SlotDraftAssignment(
        slot_id="s1",
        shape_id=_SHAPE_ID,
        shape_path=_SHAPE_PATH,
        slot_role=SlotRole.TITLE,
        semantic_label="Title",
        action=SlotDraftAction.NEEDS_INPUT,
        text=None,
        missing_information="Need the client name.",
        capacity_characters=100,
        capacity_utilization=0.0,
        max_lines_estimate=2,
    )
    draft = SlideContentDraft(
        slide_id=_SLIDE_ID,
        deck_id=_DECK_ID,
        slide_number=_SLIDE_NUMBER,
        brief_key_message="Msg.",
        assignments=[incomplete_assign],
        drafting_summary="Test.",
        is_complete=False,
    )
    req = _request(draft=draft)
    # Fails before reaching source/stable-targeting steps — no extra patches needed
    with pytest.raises(PowerPointWritebackError, match="is not complete"):
        preflight(req)


# ---------------------------------------------------------------------------
# Step 2: Draft / slot-map consistency
# ---------------------------------------------------------------------------


def test_slide_id_mismatch_rejected() -> None:
    req = _request(
        draft=_draft(slide_id="slide-OTHER"),
        slot_map=_slot_map(slide_id=_SLIDE_ID),
    )
    with pytest.raises(PowerPointWritebackError, match="slide_id mismatch"):
        preflight(req)


def test_deck_id_mismatch_rejected() -> None:
    req = _request(
        draft=_draft(deck_id="deck-A"),
        slot_map=_slot_map(deck_id="deck-B"),
    )
    with pytest.raises(PowerPointWritebackError, match="deck_id mismatch"):
        preflight(req)


def test_slide_number_mismatch_rejected() -> None:
    req = _request(
        draft=_draft(slide_number=3),
        slot_map=_slot_map(slide_number=2),
    )
    with pytest.raises(PowerPointWritebackError, match="slide_number mismatch"):
        preflight(req)


# ---------------------------------------------------------------------------
# Step 3: Per-assignment identity
# ---------------------------------------------------------------------------


def test_assignment_slot_id_not_in_slot_map_rejected() -> None:
    req = _request(
        draft=_draft(assignments=[_assignment(slot_id="unknown-slot")]),
        slot_map=_slot_map(slots=[_slot(slot_id="s1")]),
    )
    with pytest.raises(PowerPointWritebackError, match="not found in slot map"):
        preflight(req)


def test_shape_id_mismatch_rejected() -> None:
    req = _request(
        draft=_draft(assignments=[_assignment(slot_id="s1", shape_id=99)]),
        slot_map=_slot_map(slots=[_slot(slot_id="s1", shape_id=7)]),
    )
    with pytest.raises(PowerPointWritebackError, match="shape_id mismatch"):
        preflight(req)


def test_shape_path_mismatch_rejected() -> None:
    req = _request(
        draft=_draft(assignments=[_assignment(slot_id="s1", shape_id=7, shape_path="7")]),
        slot_map=_slot_map(slots=[_slot(slot_id="s1", shape_id=7, shape_path="12/7")]),
    )
    with pytest.raises(PowerPointWritebackError, match="shape_path mismatch"):
        preflight(req)


def test_slot_map_slot_with_no_assignment_rejected() -> None:
    """Slot map has two slots but draft only covers one."""
    req = _request(
        draft=_draft(assignments=[_assignment(slot_id="s1")]),
        slot_map=_slot_map(slots=[
            _slot(slot_id="s1", shape_id=7, shape_path="7"),
            _slot(slot_id="s2", shape_id=8, shape_path="8"),
        ]),
    )
    with pytest.raises(PowerPointWritebackError, match="no corresponding draft assignment"):
        preflight(req)


# ---------------------------------------------------------------------------
# Step 5: Source PPTX reachability
# ---------------------------------------------------------------------------


def test_source_pptx_not_found_rejected(tmp_path: Path) -> None:
    req = _request(source_pptx=tmp_path / "no_such.pptx")
    with pytest.raises(PowerPointWritebackError, match="not found"):
        preflight(req)


def test_slide_number_out_of_range_rejected(tmp_path: Path) -> None:
    src = tmp_path / "deck.pptx"
    src.touch()
    req = _request(
        source_pptx=src,
        slot_map=_slot_map(slide_number=99),
        draft=_draft(slide_number=99),
    )
    mock_prs = MagicMock()
    mock_prs.slides = [MagicMock(), MagicMock()]  # only 2 slides

    with (
        patch("slidestein.writeback.preflight.Presentation", return_value=mock_prs),
        # resolve_slide_by_stable_id returns ordinal 99 — range check then fails
        patch(
            "slidestein.writeback.preflight.resolve_slide_by_stable_id",
            return_value=(99, _NATIVE_SLIDE_ID),
        ),
    ):
        with pytest.raises(PowerPointWritebackError, match="out of range"):
            preflight(req)


# ---------------------------------------------------------------------------
# Step 5b: Stable slide targeting
# ---------------------------------------------------------------------------


def test_no_deck_id_rejected(tmp_path: Path) -> None:
    """slot_map without deck_id must be rejected before stable targeting."""
    src = tmp_path / "deck.pptx"
    src.touch()
    req = _request(
        source_pptx=src,
        slot_map=_slot_map(deck_id=None),
        draft=_draft(deck_id=None),
    )
    mock_prs = MagicMock()
    mock_prs.slides = [MagicMock()] * 5

    with patch("slidestein.writeback.preflight.Presentation", return_value=mock_prs):
        with pytest.raises(PowerPointWritebackError, match="deck_id is required"):
            preflight(req)


def test_stable_targeting_slide_reorder(tmp_path: Path) -> None:
    """After deck reorder, target is now at ordinal 3 (was 2). Plan must use ordinal 3."""
    src = tmp_path / "deck.pptx"
    src.touch()
    out = tmp_path / "out.pptx"
    # slot_map still records slide_number=2 (from M5.2 time)
    req = _request(source_pptx=src, output_pptx=out)

    slide = _mock_pptx_slide()
    mock_prs = MagicMock()
    mock_prs.slides = [MagicMock()] * 5
    mock_prs.slides[2] = slide  # after reorder, target is now ordinal 3 (index 2)

    with (
        patch("slidestein.writeback.preflight.Presentation", return_value=mock_prs),
        # resolve returns ordinal 3 even though slot_map says 2
        patch(
            "slidestein.writeback.preflight.resolve_slide_by_stable_id",
            return_value=(3, _NATIVE_SLIDE_ID),
        ),
        patch("slidestein.writeback.preflight.check_source_currentness"),
        patch("slidestein.identity.slide_identity.compute_structure_fingerprint", return_value="fp"),
        patch("slidestein.writeback.verification.compute_format_fingerprint", return_value="fmt"),
        patch(
            "slidestein.writeback.verification.compute_deck_structure_snapshot",
            return_value=([_NATIVE_SLIDE_ID], {str(_NATIVE_SLIDE_ID): "fp"}),
        ),
    ):
        plan = preflight(req)

    # slot_map.slide_number (ordinal at M5.2 time) is preserved
    assert plan.slide_number == _SLIDE_NUMBER
    # resolved_slide_number reflects actual current position
    assert plan.resolved_slide_number == 3
    assert plan.native_slide_id == _NATIVE_SLIDE_ID


def test_stable_targeting_not_found_rejected(tmp_path: Path) -> None:
    """resolve_slide_by_stable_id raising propagates correctly."""
    src = tmp_path / "deck.pptx"
    src.touch()
    req = _request(source_pptx=src)

    mock_prs = MagicMock()
    mock_prs.slides = [MagicMock()] * 5

    with (
        patch("slidestein.writeback.preflight.Presentation", return_value=mock_prs),
        patch(
            "slidestein.writeback.preflight.resolve_slide_by_stable_id",
            side_effect=PowerPointWritebackError("No slide resolves to slide_id"),
        ),
    ):
        with pytest.raises(PowerPointWritebackError, match="No slide resolves"):
            preflight(req)


# ---------------------------------------------------------------------------
# Step 6: Source-currentness check
# ---------------------------------------------------------------------------


def test_source_stale_rejected(tmp_path: Path) -> None:
    """check_source_currentness raising propagates as error."""
    src = tmp_path / "deck.pptx"
    src.touch()
    req = _request(source_pptx=src)

    slide = _mock_pptx_slide()
    mock_prs = MagicMock()
    mock_prs.slides = [MagicMock()] * 5
    mock_prs.slides[_SLIDE_NUMBER - 1] = slide

    with (
        patch("slidestein.writeback.preflight.Presentation", return_value=mock_prs),
        patch(
            "slidestein.writeback.preflight.resolve_slide_by_stable_id",
            return_value=(_SLIDE_NUMBER, _NATIVE_SLIDE_ID),
        ),
        patch(
            "slidestein.writeback.preflight.check_source_currentness",
            side_effect=PowerPointWritebackError("fingerprint mismatch"),
        ),
    ):
        with pytest.raises(PowerPointWritebackError, match="fingerprint mismatch"):
            preflight(req)


def test_source_currentness_passes_through_on_ok(tmp_path: Path) -> None:
    """When check_source_currentness is a no-op, preflight succeeds."""
    src = tmp_path / "deck.pptx"
    src.touch()
    out = tmp_path / "out.pptx"
    req = _request(source_pptx=src, output_pptx=out)

    plan = _patched_preflight(req)
    assert isinstance(plan, WritebackPlan)


# ---------------------------------------------------------------------------
# Step 7: Shape resolution and shape_id validation
# ---------------------------------------------------------------------------


def test_shape_path_not_in_live_source_rejected(tmp_path: Path) -> None:
    """Shape path is valid in the slot map, but not present in the PPTX."""
    src = tmp_path / "deck.pptx"
    src.touch()
    req = _request(source_pptx=src)

    from pptx.enum.shapes import MSO_SHAPE_TYPE
    shape = MagicMock()
    shape.shape_id = 99  # PPTX only has shape 99, not shape 7
    shape.shape_type = MSO_SHAPE_TYPE.AUTO_SHAPE
    shape.has_text_frame = True

    slide = MagicMock()
    slide.shapes = [shape]

    mock_prs = MagicMock()
    mock_prs.slides = [MagicMock()] * 5
    mock_prs.slides[_SLIDE_NUMBER - 1] = slide

    with (
        patch("slidestein.writeback.preflight.Presentation", return_value=mock_prs),
        patch(
            "slidestein.writeback.preflight.resolve_slide_by_stable_id",
            return_value=(_SLIDE_NUMBER, _NATIVE_SLIDE_ID),
        ),
        patch("slidestein.writeback.preflight.check_source_currentness"),
    ):
        with pytest.raises(PowerPointWritebackError, match="Cannot resolve shape_path"):
            preflight(req)


def test_shape_id_wrong_in_live_source_rejected(tmp_path: Path) -> None:
    """Slot map has shape_id=999 + shape_path='7', but PPTX shape at path '7' has id=7."""
    src = tmp_path / "deck.pptx"
    src.touch()

    req = _request(
        source_pptx=src,
        draft=_draft(assignments=[
            _assignment(slot_id="s1", shape_id=999, shape_path="7", text="x"),
        ]),
        slot_map=_slot_map(slots=[_slot(slot_id="s1", shape_id=999, shape_path="7")]),
    )

    from pptx.enum.shapes import MSO_SHAPE_TYPE
    shape = MagicMock()
    shape.shape_id = 7   # PPTX has id=7; slot map says 999
    shape.shape_type = MSO_SHAPE_TYPE.AUTO_SHAPE
    shape.has_text_frame = True

    slide = MagicMock()
    slide.shapes = [shape]

    mock_prs = MagicMock()
    mock_prs.slides = [MagicMock()] * 5
    mock_prs.slides[_SLIDE_NUMBER - 1] = slide

    with (
        patch("slidestein.writeback.preflight.Presentation", return_value=mock_prs),
        patch(
            "slidestein.writeback.preflight.resolve_slide_by_stable_id",
            return_value=(_SLIDE_NUMBER, _NATIVE_SLIDE_ID),
        ),
        patch("slidestein.writeback.preflight.check_source_currentness"),
    ):
        with pytest.raises(PowerPointWritebackError, match="shape_id=7"):
            preflight(req)


def test_shape_no_text_frame_rejected(tmp_path: Path) -> None:
    src = tmp_path / "deck.pptx"
    src.touch()
    req = _request(source_pptx=src)

    from pptx.enum.shapes import MSO_SHAPE_TYPE
    shape = MagicMock()
    shape.shape_id = _SHAPE_ID
    shape.shape_type = MSO_SHAPE_TYPE.AUTO_SHAPE
    shape.has_text_frame = False

    slide = MagicMock()
    slide.shapes = [shape]

    mock_prs = MagicMock()
    mock_prs.slides = [MagicMock()] * 5
    mock_prs.slides[_SLIDE_NUMBER - 1] = slide

    with (
        patch("slidestein.writeback.preflight.Presentation", return_value=mock_prs),
        patch(
            "slidestein.writeback.preflight.resolve_slide_by_stable_id",
            return_value=(_SLIDE_NUMBER, _NATIVE_SLIDE_ID),
        ),
        patch("slidestein.writeback.preflight.check_source_currentness"),
    ):
        with pytest.raises(PowerPointWritebackError, match="does not have a text frame"):
            preflight(req)


# ---------------------------------------------------------------------------
# Happy path — valid request produces a WritebackPlan
# ---------------------------------------------------------------------------


def test_valid_request_returns_plan(tmp_path: Path) -> None:
    src = tmp_path / "deck.pptx"
    src.touch()
    out = tmp_path / "out.pptx"
    req = _request(source_pptx=src, output_pptx=out)

    plan = _patched_preflight(req, n_slides=5)

    assert isinstance(plan, WritebackPlan)
    assert plan.slide_id == _SLIDE_ID
    assert plan.slide_number == _SLIDE_NUMBER
    assert plan.schema_version == WRITEBACK_SCHEMA_VERSION
    assert len(plan.operations) == 1
    assert plan.operations[0].slot_id == "s1"


def test_plan_has_resolved_slide_number(tmp_path: Path) -> None:
    src = tmp_path / "deck.pptx"
    src.touch()
    out = tmp_path / "out.pptx"
    req = _request(source_pptx=src, output_pptx=out)

    plan = _patched_preflight(req)
    assert plan.resolved_slide_number == _SLIDE_NUMBER


def test_plan_has_native_slide_id(tmp_path: Path) -> None:
    src = tmp_path / "deck.pptx"
    src.touch()
    out = tmp_path / "out.pptx"
    req = _request(source_pptx=src, output_pptx=out)

    plan = _patched_preflight(req)
    assert plan.native_slide_id == _NATIVE_SLIDE_ID


def test_plan_counts_replacements(tmp_path: Path) -> None:
    src = tmp_path / "deck.pptx"
    src.touch()
    out = tmp_path / "out.pptx"

    from pptx.enum.shapes import MSO_SHAPE_TYPE

    def _s(sid: str, shid: int) -> TemplateSlot:
        return _slot(slot_id=sid, shape_id=shid, shape_path=str(shid))

    def _a(sid: str, shid: int, action: SlotDraftAction, text=None) -> SlotDraftAssignment:
        return _assignment(slot_id=sid, shape_id=shid, shape_path=str(shid), action=action, text=text)

    slot_map = _slot_map(slots=[_s("s1", 7), _s("s2", 8)])
    draft = _draft(assignments=[
        _a("s1", 7, SlotDraftAction.REPLACE, "Hello"),
        _a("s2", 8, SlotDraftAction.CLEAR),
    ])
    req = _request(draft=draft, slot_map=slot_map, source_pptx=src, output_pptx=out)

    def _shape(sid: int) -> MagicMock:
        s = MagicMock()
        s.shape_id = sid
        s.shape_type = MSO_SHAPE_TYPE.AUTO_SHAPE
        s.has_text_frame = True
        return s

    slide = MagicMock()
    slide.shapes = [_shape(7), _shape(8)]

    mock_prs = MagicMock()
    mock_prs.slides = [MagicMock()] * 5
    mock_prs.slides[_SLIDE_NUMBER - 1] = slide

    with (
        patch("slidestein.writeback.preflight.Presentation", return_value=mock_prs),
        _hardening_patches(),
    ):
        plan = preflight(req)

    assert plan.replacement_count == 1
    assert plan.clear_count == 1


def test_plan_has_before_fingerprint(tmp_path: Path) -> None:
    src = tmp_path / "deck.pptx"
    src.touch()
    out = tmp_path / "out.pptx"
    req = _request(source_pptx=src, output_pptx=out)

    plan = _patched_preflight(req, fp="deadbeef")
    assert plan.before_structure_fingerprint == "deadbeef"


def test_plan_has_deck_baseline(tmp_path: Path) -> None:
    """Plan must carry the before_deck_native_ids and before_deck_structure snapshot."""
    src = tmp_path / "deck.pptx"
    src.touch()
    out = tmp_path / "out.pptx"
    req = _request(source_pptx=src, output_pptx=out)

    plan = _patched_preflight(req, fp="aabbccdd")
    assert plan.before_deck_native_ids == [_NATIVE_SLIDE_ID]
    assert plan.before_deck_structure == {str(_NATIVE_SLIDE_ID): "aabbccdd"}


# ---------------------------------------------------------------------------
# Output path safety (additional)
# ---------------------------------------------------------------------------


def test_output_exists_without_overwrite_rejected(tmp_path: Path) -> None:
    src = tmp_path / "deck.pptx"
    src.touch()
    out = tmp_path / "out.pptx"
    out.touch()  # exists already

    req = _request(source_pptx=src, output_pptx=out, overwrite=False)
    mock_prs = MagicMock()
    mock_prs.slides = [MagicMock()] * 5
    with patch("slidestein.writeback.preflight.Presentation", return_value=mock_prs):
        with pytest.raises(PowerPointWritebackError, match="already exists"):
            preflight(req)
