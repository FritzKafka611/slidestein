"""Shared pytest fixtures."""

from __future__ import annotations

import pytest

from slidestein.config import Settings


@pytest.fixture
def settings(tmp_path) -> Settings:
    return Settings(
        slides_dir=tmp_path / "slides",
        output_dir=tmp_path / "output",
        previews_dir=tmp_path / "previews",
        db_path=tmp_path / "slidestein.db",
        lancedb_uri=str(tmp_path / "lancedb"),
    )


@pytest.fixture
def three_slide_pptx(tmp_path) -> "Path":
    """A minimal 3-slide PPTX with known text content."""
    from pathlib import Path

    from pptx import Presentation

    prs = Presentation()
    layout = prs.slide_layouts[0]  # Title Slide: idx-0 = title, idx-1 = subtitle

    for i in range(1, 4):
        slide = prs.slides.add_slide(layout)
        for ph in slide.placeholders:
            if ph.placeholder_format.idx == 0:
                ph.text = f"Slide {i} Title"
            elif ph.placeholder_format.idx == 1:
                ph.text = f"Subtitle for slide {i}"

    path = tmp_path / "fixture.pptx"
    prs.save(str(path))
    return path
