"""M3.4.1 empirical identity validation.

Each experiment uses a SEPARATE temporary directory with its own copy of the
deck at a fixed path so that re-indexing triggers lineage rule B
(same path, new fingerprint → preserve deck_id).

Never modifies the original sample deck.
"""
from __future__ import annotations

import io
import shutil
import sys
import tempfile
import zipfile
from pathlib import Path

from lxml import etree
from pptx import Presentation

sys.path.insert(0, str(Path(__file__).parents[1] / "src"))

from slidestein.ingestion.service import IngestionService
from slidestein.library.store import SlideLibrary
from slidestein.pptx.python_pptx_adapter import PythonPptxAdapter

SAMPLE = Path("data/test_decks/sample_consulting_deck.pptx")

_NS_P  = "http://schemas.openxmlformats.org/presentationml/2006/main"
_NS_R  = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"


# ---------------------------------------------------------------------------
# PPTX manipulation helpers (all overwrite dest in-place from src bytes)
# ---------------------------------------------------------------------------

def _reorder_first_to_last(src: Path, dst: Path) -> None:
    """Copy src to dst with first slide moved to last position."""
    with zipfile.ZipFile(src, "r") as z:
        entries = {n: z.read(n) for n in z.namelist()}
    root = etree.fromstring(entries["ppt/presentation.xml"])
    lst  = root.find(f"{{{_NS_P}}}sldIdLst")
    first = lst[0]; lst.remove(first); lst.append(first)
    entries["ppt/presentation.xml"] = etree.tostring(
        root, xml_declaration=True, encoding="UTF-8", standalone=True
    )
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        for n, d in entries.items():
            z.writestr(n, d)
    dst.write_bytes(buf.getvalue())


def _remove_last_slide(src: Path, dst: Path) -> None:
    """Copy src to dst with last slide removed."""
    _NS_RL = "http://schemas.openxmlformats.org/package/2006/relationships"
    with zipfile.ZipFile(src, "r") as z:
        entries = {n: z.read(n) for n in z.namelist()}
    prs_root  = etree.fromstring(entries["ppt/presentation.xml"])
    rels_root = etree.fromstring(entries["ppt/_rels/presentation.xml.rels"])
    lst  = prs_root.find(f"{{{_NS_P}}}sldIdLst")
    last = lst[-1]
    r_id = last.get(f"{{{_NS_R}}}id")
    slide_target = None
    for rel in list(rels_root):
        if rel.get("Id") == r_id:
            slide_target = rel.get("Target")
            rels_root.remove(rel)
            break
    lst.remove(last)
    entries["ppt/presentation.xml"] = etree.tostring(
        prs_root, xml_declaration=True, encoding="UTF-8", standalone=True
    )
    entries["ppt/_rels/presentation.xml.rels"] = etree.tostring(
        rels_root, xml_declaration=True, encoding="UTF-8", standalone=True
    )
    if slide_target:
        fname = slide_target.rsplit("/", 1)[-1]
        entries.pop(f"ppt/{slide_target}", None)
        entries.pop(f"ppt/slides/_rels/{fname}.rels", None)
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        for n, d in entries.items():
            z.writestr(n, d)
    dst.write_bytes(buf.getvalue())


# ---------------------------------------------------------------------------
# Ingestion helper
# ---------------------------------------------------------------------------

def _ingest(db_path: Path, pptx_path: Path) -> list:
    lib = SlideLibrary(db_path, str(db_path.with_suffix(".lance")))
    lib.init()
    svc = IngestionService(lib, PythonPptxAdapter(), db_path.parent / "prev")
    records = svc.ingest_deck(pptx_path, render_previews=False)
    lib.close()
    return records


def _fmt(records, title: str = "", limit: int | None = None) -> None:
    if title:
        print(f"\n=== {title} ===")
    print(f"  {'native_id':>12}  {'slide_id':>38}  {'#':>3}  {'content_fp[:10]':>12}  {'struct_fp[:10]':>12}  act")
    shown = sorted(records, key=lambda r: r.native_slide_id)
    if limit:
        shown = shown[:limit]
    for r in shown:
        cfp = (r.content_fingerprint or "")[:10]
        sfp = (r.structure_fingerprint or "")[:10]
        print(
            f"  {r.native_slide_id:>12}  {r.slide_id}  {r.slide_number:>3}  {cfp:>12}  {sfp:>12}  {r.is_active}"
        )


# ---------------------------------------------------------------------------
# Experiment runner
# ---------------------------------------------------------------------------

def _experiment(name: str, tmp_root: Path, modify_fn) -> tuple[list, list]:
    """Ingest baseline copy, apply modify_fn in-place, re-ingest, return (before, after)."""
    exp_dir = tmp_root / name
    exp_dir.mkdir()
    pptx = exp_dir / "deck.pptx"
    db   = exp_dir / "lib.db"
    shutil.copy(SAMPLE, pptx)

    before = _ingest(db, pptx)
    modify_fn(pptx)                  # modifies pptx in-place
    after  = _ingest(db, pptx)

    return before, after


# ---------------------------------------------------------------------------
# Main
# ---------------------------------------------------------------------------

def main() -> None:
    assert SAMPLE.exists(), f"Source deck not found: {SAMPLE}"

    with tempfile.TemporaryDirectory() as _tmp:
        tmp = Path(_tmp)

        # ── A. BASELINE ─────────────────────────────────────────────────────
        exp_a = tmp / "A"
        exp_a.mkdir()
        pptx_a = exp_a / "deck.pptx"
        db_a   = exp_a / "lib.db"
        shutil.copy(SAMPLE, pptx_a)
        rec_a = _ingest(db_a, pptx_a)
        _fmt(rec_a, "A. Baseline (16 slides)")
        deck_id_a      = rec_a[0].deck_id
        id_by_native_a = {r.native_slide_id: r.slide_id for r in rec_a}
        print(f"  deck_id = {deck_id_a}")

        # ── B. REORDER: first slide → last ──────────────────────────────────
        def _mod_b(pptx: Path) -> None:
            _reorder_first_to_last(pptx, pptx)   # src == dst is safe (reads all first)

        before_b, after_b = _experiment("B", tmp, _mod_b)
        first_native = before_b[0].native_slide_id
        first_slide_id_before = before_b[0].slide_id
        moved = next(r for r in after_b if r.native_slide_id == first_native)
        _fmt(after_b, "B. After reorder (first -> last)")
        print(f"  deck_id unchanged    = {after_b[0].deck_id == before_b[0].deck_id}")
        print(f"  native_id of moved   = {moved.native_slide_id}  (was slide_number=1)")
        print(f"  slide_id unchanged   = {moved.slide_id == first_slide_id_before}")
        print(f"  slide_number after   = {moved.slide_number}  (was 1, expected last)")

        # ── C. TEXT EDIT on slide 1 ─────────────────────────────────────────
        def _mod_c(pptx: Path) -> None:
            prs = Presentation(str(pptx))
            for ph in prs.slides[0].placeholders:
                if ph.placeholder_format.idx == 0:
                    ph.text = "EDITED TITLE XYZ"
                    break
            prs.save(str(pptx))

        before_c, after_c = _experiment("C", tmp, _mod_c)
        s1_before = next(r for r in before_c if r.slide_number == 1)
        s1_after  = next(r for r in after_c  if r.native_slide_id == s1_before.native_slide_id)
        _fmt(after_c[:3], "C. Text edit on slide 1 (first 3 by native_id)")
        print(f"  deck_id unchanged    = {after_c[0].deck_id == before_c[0].deck_id}")
        print(f"  slide_id unchanged   = {s1_after.slide_id == s1_before.slide_id}")
        print(f"  content_fp changed   = {s1_after.content_fingerprint != s1_before.content_fingerprint}")
        print(f"  struct_fp unchanged  = {s1_after.structure_fingerprint == s1_before.structure_fingerprint}")

        # ── D. INSERT new slide at end ──────────────────────────────────────
        def _mod_d(pptx: Path) -> None:
            prs = Presentation(str(pptx))
            new_s = prs.slides.add_slide(prs.slide_layouts[0])
            new_s.placeholders[0].text = "NEW INSERTED SLIDE"
            prs.save(str(pptx))

        before_d, after_d = _experiment("D", tmp, _mod_d)
        id_by_native_d = {r.native_slide_id: r.slide_id for r in before_d}
        all_unchanged = all(
            r.slide_id == id_by_native_d[r.native_slide_id]
            for r in after_d if r.native_slide_id in id_by_native_d
        )
        new_slides = [r for r in after_d if r.native_slide_id not in id_by_native_d]
        _fmt(after_d[-3:], "D. Insert slide at end (last 3 by native_id)")
        print(f"  deck_id unchanged               = {after_d[0].deck_id == before_d[0].deck_id}")
        print(f"  all existing slide_ids unchanged = {all_unchanged}")
        print(f"  new slide count                  = {len(new_slides)}")
        if new_slides:
            ns = new_slides[0]
            print(f"  new slide_id                     = {ns.slide_id}  slide_number={ns.slide_number}")

        # ── E. DELETE last slide ─────────────────────────────────────────────
        def _mod_e(pptx: Path) -> None:
            _remove_last_slide(pptx, pptx)

        before_e, after_e = _experiment("E", tmp, _mod_e)
        exp_e_dir = tmp / "E"
        lib_e = SlideLibrary(exp_e_dir / "lib.db", str((exp_e_dir / "lib").with_suffix(".lance")))
        lib_e.init()
        removed_native    = before_e[-1].native_slide_id
        removed_slide_id  = before_e[-1].slide_id
        removed_rec       = lib_e.get_record(removed_slide_id)
        active_after_e    = lib_e.list_active_slides()
        lib_e.close()
        _fmt(after_e, "E. After deleting last slide (15 remain)")
        print(f"  deck_id unchanged            = {after_e[0].deck_id == before_e[0].deck_id}")
        print(f"  removed slide in DB          = {removed_rec is not None}")
        print(f"  removed slide is_active      = {removed_rec.is_active if removed_rec else 'N/A'}")
        print(f"  active slide count           = {len(active_after_e)}")

    print("\n[OK] Empirical validation complete — no Anthropic API calls made.")


if __name__ == "__main__":
    main()
